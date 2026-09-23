"""
Phase 7 validation - the rendered dashboard vs Excel.

Opens outputs/<TICKER>/dashboard/index.html in headless Chromium and, for every scenario x page:
  * reads every element marked [data-ctl] (headline totals the page computed from the exported records)
  * compares its unrounded value (data-raw) with the Excel control total, and its DISPLAYED text with the
    Excel value rounded to 0.1 ($m)
  * Base scenario is also compared with the control totals read directly from the saved Excel workbook
Also checks: no JavaScript errors, every page renders, no horizontal page scroll at 400px width.
Writes outputs/<TICKER>/tests/phase7_dashboard_validation.json (+ screenshots in outputs/<TICKER>/tests/screens/). Exit 1 on any failure.
Run: python tests/validate_dashboard.py --company TICKER   (default MCHP)
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import engine  # noqa: E402
EN = engine.get()
HTML = EN.index_html.resolve()
DATA = json.loads(EN.data_json.read_text())
SCREENS = EN.screens
SCREENS.mkdir(exist_ok=True)
SCENARIOS = ["Downside", "Base", "Upside"]
PAGES = ["overview", "operations", "variance", "scenarios", "capital", "validation"]


def excel_controls():
    ws = load_workbook(EN.model, data_only=True)["Dash_Export"]
    r = next(i for i in range(1, ws.max_row + 1) if ws.cell(row=i, column=1).value == "TABLE: control_totals")
    out, i = {}, r + 2
    while ws.cell(row=i, column=1).value:
        out[ws.cell(row=i, column=1).value] = ws.cell(row=i, column=2).value
        i += 1
    scen = ws.cell(row=next(i for i in range(1, ws.max_row + 1) if ws.cell(row=i, column=1).value == "scenario"), column=2).value
    return scen, out


def parse_display(t):
    t = t.replace("−", "-").replace("$", "").replace(",", "").replace("m", "").strip()
    return float(re.sub(r"^\+", "", t))


def main():
    xl_scen, xl = excel_controls()
    results, errors = [], []
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page(viewport={"width": 1360, "height": 900})
        pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.on("console", lambda m: errors.append(m.text) if m.type == "error" and "fonts.g" not in m.text
              and "ERR_" not in m.text else None)
        pg.goto(HTML.as_uri())
        pg.wait_for_function("window.__ECD__ !== undefined")
        seen = set()
        for s in SCENARIOS:
            ctl = {x["control_id"]: x["value"] for x in DATA["scenarios"][s]["control_totals"]}
            for page in PAGES:
                pg.evaluate(f"window.__ECD__.setScenario('{s}'); window.__ECD__.setPage('{page}')")
                visible = pg.evaluate(f"!document.querySelector('#p-{page}').hidden && document.querySelector('#p-{page}').innerText.length")
                results.append({"check": f"{s}/{page}: page renders", "status": "PASS" if visible and visible > 200 else "FAIL"})
                for e in pg.query_selector_all(f"#p-{page} [data-ctl]"):
                    cid, raw, text = e.get_attribute("data-ctl"), float(e.get_attribute("data-raw")), e.text_content()
                    exp = ctl[cid]
                    ok_raw = abs(raw - exp) <= 1e-6 * max(1, abs(exp))
                    ok_txt = abs(parse_display(text) - round(exp, 1)) < 1e-9
                    results.append({"check": f"{s}/{page}: {cid}", "excel": exp, "dashboard_raw": raw, "displayed": text,
                                    "status": "PASS" if ok_raw and ok_txt else "FAIL"})
                    seen.add((s, cid))
                    if s == xl_scen and cid in xl:
                        ok = abs(raw - xl[cid]) <= 1e-6 * max(1, abs(xl[cid]))
                        results.append({"check": f"{s}/{page}: {cid} vs workbook cell (direct)", "excel": xl[cid],
                                        "dashboard_raw": raw, "status": "PASS" if ok else "FAIL"})
            # Validation page: every in-page comparison must pass
            pg.evaluate("window.__ECD__.setPage('validation')")
            chip = pg.inner_text("#valChip")
            results.append({"check": f"{s}: in-page validation table", "displayed": chip,
                            "status": "PASS" if chip.startswith("30 / 30") else "FAIL"})
        missing = [(s, c) for s in SCENARIOS for c in DATA["scenarios"][s][ "control_totals"] and
                   [x["control_id"] for x in DATA["scenarios"][s]["control_totals"]] if (s, c) not in seen
                   and not (c == "capex_npv_base" and s != "Base")]
        results.append({"check": "Every control total is displayed somewhere in every scenario", "missing": missing,
                        "status": "PASS" if not missing else "FAIL"})
        # screenshots (Base, dark + light)
        pg.evaluate("window.__ECD__.setScenario('Base')")
        for page in ["overview", "operations", "variance", "capital"]:
            pg.evaluate(f"window.__ECD__.setPage('{page}')")
            pg.screenshot(path=str(SCREENS / f"{page}_dark.png"), full_page=True)
        pg.evaluate("document.documentElement.dataset.theme='light'; window.__ECD__.setPage('overview')")
        pg.wait_for_timeout(100)
        pg.screenshot(path=str(SCREENS / "overview_light.png"), full_page=True)
        # phone width: no horizontal page scroll
        m = b.new_page(viewport={"width": 400, "height": 860})
        m.on("pageerror", lambda e: errors.append(str(e)))
        m.goto(HTML.as_uri())
        m.wait_for_function("window.__ECD__ !== undefined")
        for page in PAGES:
            m.evaluate(f"window.__ECD__.setPage('{page}')")
            sw = m.evaluate("document.documentElement.scrollWidth")
            results.append({"check": f"400px width/{page}: no horizontal page scroll", "scroll_width": sw,
                            "status": "PASS" if sw <= 401 else "FAIL"})
        m.evaluate("window.__ECD__.setPage('overview')")
        m.screenshot(path=str(SCREENS / "overview_phone.png"), full_page=True)
        b.close()
    results.append({"check": "No JavaScript errors", "errors": errors[:5], "status": "PASS" if not errors else "FAIL"})
    compared = sum(1 for r in results if "dashboard_raw" in r)
    failed = sum(1 for r in results if r["status"] == "FAIL")
    out = {"run_utc": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"), "data_sha256": DATA["source_sha256"],
           "compared": compared, "failed": failed, "total_checks": len(results), "scenarios": SCENARIOS, "results": results}
    (EN.tests / "phase7_dashboard_validation.json").write_text(json.dumps(out, indent=2))
    print(f"Dashboard validation: {len(results) - failed} PASS, {failed} FAIL; displayed values compared with Excel: {compared}")
    for r in results:
        if r["status"] == "FAIL":
            print("  FAIL:", r)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()

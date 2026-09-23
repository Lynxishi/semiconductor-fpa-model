"""
Phase 2 validation - scenario selection and operating-model logic.

For each scenario (Downside, Base, Upside):
  1. copy the workbook, set the ONE selector cell (Assumptions!C6), recalculate with LibreOffice
  2. rebuild the operating model independently in Python from the raw inputs on the Assumptions tab
  3. compare every month x product for units sold, revenue, cost of revenue, gross profit and closing inventory
  4. read the Checks tab status

Run: python tests/validate_phase2_scenarios.py --company TICKER
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import EN, PCOL, PRODUCTS, BY, base_mix, find_row  # noqa: E402,F401

ROOT = Path(__file__).resolve().parents[1]
SRC = EN.model
RECALC = str(Path(__file__).resolve().parents[1] / "scripts" / "recalc.py")
TMP = EN.tests / "_scenario_runs"


def python_model(a):
    """Independent re-implementation. a = raw inputs read from the Assumptions tab (hardcoded cells only)."""
    idx = {"Downside": 0, "Base": 1, "Upside": 2}[a["scenario"]]
    drv = {k: v[idx] for k, v in a["drivers"].items()}
    out = {}
    for p in PRODUCTS:
        x = a["prod"][p]
        rev0 = a["base_rev"] * x["mix"]
        sold0 = rev0 / x["asp0"]
        prod0 = sold0 / x["yld0"]
        cogs0 = prod0 * x["vc0"] + x["fix0"] * 12
        cpgu0 = cogs0 / sold0
        dem_fy = sold0 * (1 + drv["dem_g"] + x["gdiff"])
        asp = x["asp0"] * (1 + drv["asp_chg"])
        yld = min(1, max(0, x["yld0"] + drv["yld_adj"]))
        vcu = x["vc0"] * (1 + drv["mat_inf"])
        inv_u = x["inv_m"] * sold0 / 12
        inv_v = inv_u * cpgu0
        months = []
        for s in a["season"]:
            dem = dem_fy * s / 12
            produced = x["cap0"] * drv["util"]
            good = produced * yld
            avail = inv_u + good
            sold = min(dem, avail)
            mfg = produced * vcu + x["fix0"]
            avg = (inv_v + mfg) / avail
            cor = sold * avg
            inv_u, inv_v = avail - sold, inv_v + mfg - cor
            months.append(dict(sold=sold, rev=sold * asp, cor=cor, gp=sold * asp - cor, inv_close=inv_u))
        out[p] = months
    return out


def read_inputs(ws, scenario):
    a = {"scenario": scenario, "drivers": {}, "prod": {p: {} for p in PRODUCTS}}
    labels = {"dem_g": "Customer demand growth", "asp_chg": "Average selling price change",
              "yld_adj": "Manufacturing yield change (percentage points vs base-year yield)",
              "util": "Capacity utilization (internal final test)",
              "mat_inf": "Material-cost inflation (variable cost per unit)"}
    for k, lab in labels.items():
        r = find_row(ws, lab)
        a["drivers"][k] = [ws.cell(row=r, column=c).value for c in (3, 4, 5)]
    a["base_rev"] = ws.cell(row=find_row(ws, f"Division net revenue, {BY} base year"), column=6).value
    plabels = {"asp0": "Base-year average selling price", "yld0": "Base-year final-test yield",
               "vc0": "Base-year variable cost per unit produced", "fix0": "Fixed manufacturing cost per month",
               "cap0": "Internal final-test capacity", "gdiff": "Product demand growth vs division driver",
               "inv_m": "Opening finished-goods inventory (months of base-year sales)"}
    for k, lab in plabels.items():
        r = find_row(ws, lab)
        for p in PRODUCTS:
            a["prod"][p][k] = ws.cell(row=r, column=PCOL[p]).value
    return a


def main():
    TMP.mkdir(exist_ok=True)
    # Reported mix read directly from the reported CSV-derived Hist_Data values (independent of Hist_Analysis)
    base = load_workbook(SRC, data_only=True)
    hd = base["Hist_Data"]
    mix = base_mix(hd)
    all_ok, summary = True, []
    for scen in ["Downside", "Base", "Upside"]:
        f = TMP / f"model_{scen}.xlsx"
        shutil.copy(SRC, f)
        wb = load_workbook(f)
        wb["Assumptions"]["C6"] = scen
        wb.save(f)
        res = json.loads(subprocess.run([sys.executable, RECALC, str(f), "90"], capture_output=True, text=True).stdout)
        wbv = load_workbook(f, data_only=True)
        asm, ops, ck = wbv["Assumptions"], wbv["Ops_Model"], wbv["Checks"]
        a = read_inputs(asm, scen)
        season_r = find_row(asm, "Seasonality factor (1.00 = average month)")
        a["season"] = [asm.cell(row=season_r, column=c).value for c in range(3, 15)]
        for p in PRODUCTS:
            a["prod"][p]["mix"] = mix[p]
        py = python_model(a)

        # locate product blocks on Ops_Model
        heads = {p["key"]: p["name"] for p in EN.products}
        max_diff, n = 0.0, 0
        tot = {"rev": 0.0, "gp": 0.0}
        for p in PRODUCTS:
            start = next(r for r in range(1, ops.max_row + 1)
                         if str(ops.cell(row=r, column=1).value or "").startswith(heads[p] + "  -  "))
            rows = {k: find_row(ops, lab, start=start) for k, lab in
                    [("sold", "Units sold"), ("rev", "Revenue"), ("cor", "Cost of revenue"),
                     ("gp", "Gross profit"), ("inv_close", "Closing finished-goods inventory")]}
            for m in range(12):
                for k, rr in rows.items():
                    xl = ops.cell(row=rr, column=3 + m).value
                    d = abs(xl - py[p][m][k])
                    max_diff = max(max_diff, d)
                    n += 1
            tot["rev"] += sum(x["rev"] for x in py[p])
            tot["gp"] += sum(x["gp"] for x in py[p])
        div_start = next(r for r in range(1, ops.max_row + 1)
                         if str(ops.cell(row=r, column=1).value or "").startswith("Division total"))
        xl_rev = ops.cell(row=find_row(ops, "Revenue", start=div_start), column=15).value
        xl_gp = ops.cell(row=find_row(ops, "Gross profit", start=div_start), column=15).value
        status_r = find_row(ck, "OVERALL MODEL STATUS")
        overall = ck.cell(row=status_r, column=3).value
        shown = ops["B4"].value
        ok = (max_diff < 1e-6 and abs(xl_rev - tot["rev"]) < 1e-6 and abs(xl_gp - tot["gp"]) < 1e-6
              and overall == "ALL CHECKS PASS" and shown == scen and res.get("total_errors") == 0)
        all_ok &= ok
        summary.append(dict(scenario=scen, excel_revenue=round(xl_rev, 3), python_revenue=round(tot["rev"], 3),
                            excel_gm=round(xl_gp / xl_rev, 4), cells_compared=n, max_abs_diff=max_diff,
                            formula_errors=res.get("total_errors"), checks_tab=overall, result="PASS" if ok else "FAIL"))
    for s in summary:
        print(s)
    (EN.tests / "phase2_scenario_results.json").write_text(json.dumps(summary, indent=2))
    shutil.rmtree(TMP, ignore_errors=True)  # Windows/OneDrive can briefly lock the temp copies
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()

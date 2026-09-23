"""
Build, test and publish one company end to end, and write outputs/<TICKER>/tests/VALIDATION_SUMMARY.md.

    python scripts/run_all_validations.py --company MCHP
    python scripts/run_all_validations.py --company TXN

Order: config check -> reported data -> model build + recalc -> Phase 1..5 independent tests -> dashboard export +
JSON validation -> dashboard page + browser test -> deck -> docs (-> regression vs the saved baseline, MCHP only).
Every step reports PASS or FAIL; the run stops being "ALL PASS" if any step fails.
"""
import datetime as dt
import json
import os
import subprocess
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import engine  # noqa: E402

ROOT = engine.ROOT


def steps(EN):
    m = str(EN.model.relative_to(ROOT))
    s = [
        ("Reported data (SEC filings) internal consistency", ["tests/validate_reported_data.py"]),
        ("Build Excel model", ["scripts/build_model.py"]),
        ("Recalculate Excel model (formula errors)", ["scripts/recalc.py", m, "180"]),
        ("Phase 1 - historical analysis: Python vs Excel", ["tests/validate_phase1_model.py"]),
        ("Phase 2 - operating model, all 3 scenarios", ["tests/validate_phase2_scenarios.py"]),
        ("Phase 3 - revenue, P&L, FCF, all 3 scenarios", ["tests/validate_phase3_scenarios.py"]),
        ("Phase 4 - budget vs actual & variance (Base and Downside selector)", ["tests/validate_phase4_variance.py"]),
        ("Phase 5 - capex NPV / IRR / payback / break-even", ["tests/validate_phase5_capex.py"]),
        ("Phase 6 - dashboard export (3 scenarios) + JSON vs Excel", ["scripts/export_dashboard_data.py"]),
        ("Phase 7 - build dashboard page", ["scripts/build_dashboard.py"]),
        ("Phase 7 - rendered dashboard vs Excel (headless browser)", ["tests/validate_dashboard.py"]),
        ("Phase 7 - rebuild dashboard with test result embedded", ["scripts/build_dashboard.py"]),
        ("Phase 8 - collect presentation figures and wording from the validated export", ["scripts/prepare_deck_data.py"]),
        ("Phase 8 - build executive presentation (.pptx)", ["node", "scripts/build_presentation.js"]),
        ("Phase 8 - extract deck text for review", ["-m", "markitdown", str(EN.deck.relative_to(ROOT)),
                                                     "-o", str((EN.deck.parent / "deck_text.md").relative_to(ROOT))]),
        ("Phase 9 - regenerate sources & assumptions document", ["scripts/build_docs.py"]),
    ]
    if EN.ticker == "MCHP" and (ROOT / "tests" / "baseline_mchp" / "data.json").exists():
        s.append(("Phase 10 - regression: Microchip outputs unchanged vs pre-engine baseline", ["tests/regress_baseline.py"]))
    return s


def main():
    EN = engine.get()
    env = dict(os.environ, FPA_COMPANY=EN.ticker)
    rows, ok_all = [], True
    print(f"Company: {EN.ticker} - {EN.company['name']} / {EN.division['name']} / {EN.fy}")
    for label, cmd in steps(EN):
        t0 = time.time()
        argv = (cmd + ["--company", EN.ticker]) if cmd[0] == "node" else [sys.executable] + cmd
        argv = [str(ROOT / c) if (c.endswith(".py") or c.endswith(".js")) and "/" in c and not c.startswith("/") else c for c in argv]
        p = subprocess.run(argv, cwd=ROOT, capture_output=True, text=True, timeout=1200, env=env)
        ok = p.returncode == 0
        tail = (p.stdout.strip().splitlines() or p.stderr.strip().splitlines() or [""])[-1][:160]
        if cmd[0].endswith("recalc.py"):
            try:
                res = json.loads(p.stdout)
                ok = res.get("status") == "success" and res.get("total_errors") == 0
                tail = f"{res.get('total_formulas')} formulas, {res.get('total_errors')} errors"
            except json.JSONDecodeError:
                ok, tail = False, (p.stderr.strip().splitlines() or ["recalc failed"])[-1][:160]
        if cmd[0] == "-m":
            tail = f"wrote {Path(cmd[-1]).name}" if ok else tail
        ok_all &= ok
        rows.append((label, "PASS" if ok else "FAIL", f"{time.time() - t0:.0f}s", tail))
        print(f"{'PASS' if ok else 'FAIL'}  {label}  ({tail})")
        if not ok:
            err = (p.stderr or "").strip()
            if err:
                print("      " + "\n      ".join(err.splitlines()[-8:]))
            if "markitdown" in cmd and "pptx" in err.lower():
                print('      -> fix: pip install "markitdown[pptx]"')
            if label.startswith(("Build Excel", "Recalculate")):
                break
    n_pass = n_fail = 0
    if EN.model.exists():
        ck = load_workbook(EN.model, data_only=True)["Checks"]
        statuses = [ck.cell(row=r, column=9).value for r in range(1, ck.max_row + 1)]
        n_pass, n_fail = statuses.count("PASS"), statuses.count("FAIL")
    lines = [f"# Validation summary - {EN.company['name']} ({EN.ticker})", "",
             f"Run: {dt.datetime.now(dt.timezone.utc).isoformat(timespec='seconds')} UTC  |  Division: {EN.division['name']}  |  "
             f"Forecast year: {EN.fy}", "",
             f"**Overall: {'ALL PASS' if ok_all and n_fail == 0 else 'FAILURES'}**  |  Excel Checks tab: {n_pass} pass / {n_fail} fail", "",
             "| Step | Result | Time | Detail |", "|---|---|---|---|"]
    lines += [f"| {a} | {b} | {c} | {d.replace('|', '/')} |" for a, b, c, d in rows]
    lines += ["", "Dashboard totals vs Excel: the JSON export is checked against Excel (Phase 6) and the rendered page's "
                  "displayed totals are checked against the Excel control totals in a headless browser (Phase 7).", "",
              f"Outputs: `{EN.index_html.relative_to(ROOT)}`, `{EN.model.relative_to(ROOT)}`, `{EN.deck.relative_to(ROOT)}`"]
    (EN.tests / "VALIDATION_SUMMARY.md").write_text("\n".join(lines) + "\n")
    print(f"\n{'ALL PASS' if ok_all and n_fail == 0 else 'FAILURES'} - summary: {(EN.tests / 'VALIDATION_SUMMARY.md').relative_to(ROOT)}")
    sys.exit(0 if ok_all and n_fail == 0 else 1)


if __name__ == "__main__":
    main()

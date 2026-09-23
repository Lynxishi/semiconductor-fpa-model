"""
Phase 5 validation - independent Python rebuild of the capex decision for all three scenarios.

Reads RAW inputs (Capex section A, Assumptions scenario columns), rebuilds the with/without cash flows,
computes NPV, IRR (bisection), payback, break-even utilization, then:
  * compares with Excel
  * re-runs the cash flows with incremental units REPLACED by break-even utilization x full added units
    and confirms NPV = 0 (independent proof of the break-even figure)
Run: python tests/validate_phase5_capex.py   (uses the saved, recalculated workbook)
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_phase2_scenarios import find_row  # noqa: E402
from _common import EN, BY, FY, PCOL, base_mix  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
wb = load_workbook(EN.model, data_only=True)
CX = EN.cfg["capex"]
TL = CX.get("target_label", EN.target["id"])
TC = PCOL[EN.target["key"]]   # target product column on Assumptions
cx, asm = wb["Capex"], wb["Assumptions"]
F = lambda lab: cx.cell(row=find_row(cx, lab), column=6).value  # noqa: E731
P = {"price": F("Equipment purchase price"), "install": F("Installation, qualification and software"),
     "add": F("Additional final-test capacity"), "yg": F(f"Expected {TL} yield improvement"),
     "sav": F("Operating-cost savings"), "mnt": F("Maintenance cost of new equipment"),
     "life": F("Useful life (and tax depreciation period)"), "salv": F("Salvage value at end of life"),
     "umax": F("Maximum practical utilization"), "ero": F(f"Annual {TL} ASP erosion after {FY}"),
     "r": F("Discount rate (division cost of capital)"), "hurdle": F("Decision criterion: minimum IRR (hurdle)"),
     "maxpb": F("Decision criterion: maximum payback"), "dso": asm.cell(row=find_row(asm, "Days sales outstanding (DSO)"), column=6).value}
A = lambda lab, col: asm.cell(row=find_row(asm, lab), column=col).value  # noqa: E731
hd = wb["Hist_Data"]
fpga_mix = base_mix(hd)[EN.target["key"]]   # target-product share of base-year revenue
base_rev = A(f"Division net revenue, {BY} base year", 6)
lg_row = find_row(cx, f"{TL} demand growth per year after {FY}")


def irr(cf):
    f = lambda x: sum(c / (1 + x) ** t for t, c in enumerate(cf))  # noqa: E731
    lo, hi = -0.99, 10.0
    if f(lo) * f(hi) > 0:
        return None
    for _ in range(300):
        m = (lo + hi) / 2
        lo, hi = (lo, m) if f(lo) * f(m) <= 0 else (m, hi)
    return (lo + hi) / 2


def project(col, units_override=None):
    asp0, vc0, y0b, cap0 = A("Base-year average selling price", TC), A("Base-year variable cost per unit produced", TC), \
        A("Base-year final-test yield", TC), A("Internal final-test capacity", TC)
    gdiff = A("Product demand growth vs division driver", TC)
    sold0 = base_rev * fpga_mix / asp0
    dem27 = sold0 * (1 + A("Customer demand growth", col) + gdiff)
    asp27 = asp0 * (1 + A("Average selling price change", col))
    vc = vc0 * (1 + A("Material-cost inflation (variable cost per unit)", col))
    y_old = min(1, max(0, y0b + A("Manufacturing yield change (percentage points vs base-year yield)", col)))
    y_new = min(1, y_old + P["yg"])
    tax = A("Tax rate", col)
    lg = cx.cell(row=lg_row, column=col).value
    cap_wo, cap_wi = cap0 * 12 * P["umax"] * y_old, (cap0 + P["add"]) * 12 * P["umax"] * y_new
    full = P["add"] * 12 * y_new
    capex = P["price"] + P["install"]
    cf, wc_prev, d_units_list = [-capex], 0.0, []
    for t in range(1, P["life"] + 1):
        dem, asp = dem27 * (1 + lg) ** t, asp27 * (1 - P["ero"]) ** t
        wo = min(dem, cap_wo)
        du = (min(dem, cap_wi) - wo) if units_override is None else units_override * full
        d_units_list.append(du)
        ebitda = du * (asp - vc / y_new) + wo * (1 / y_old - 1 / y_new) * vc + P["sav"] - P["mnt"]
        dep = (capex - P["salv"]) / P["life"]
        nopat = (ebitda - dep) * (1 - tax)
        wc = 0 if t == P["life"] else du * asp * P["dso"] / 365
        cf.append(nopat + dep - (wc - wc_prev) + (P["salv"] if t == P["life"] else 0))
        wc_prev = wc
    npv = sum(c / (1 + P["r"]) ** t for t, c in enumerate(cf))
    cum, pb = 0.0, None
    for t, c in enumerate(cf):
        if t and cum < 0 <= cum + c:
            pb = t - 1 + (-cum) / c
        cum += c
    return dict(npv=npv, irr=irr(cf), pb=pb, cf=cf, full=full, d_units=d_units_list)


out, ok_all = [], True
for name, col in [("Downside", 3), ("Base", 4), ("Upside", 5)]:
    py = project(col)
    # break-even: NPV is linear in incremental units -> solve with two evaluations
    n0, n1 = project(col, 0.0)["npv"], project(col, 1.0)["npv"]
    be = 0.0 if n0 >= 0 else -n0 / (n1 - n0)
    npv_at_be = project(col, be)["npv"]
    blk = next(r for r in range(1, cx.max_row + 1) if str(cx.cell(row=r, column=1).value or "").startswith(f"C") and
               name.upper() in str(cx.cell(row=r, column=1).value))
    X = lambda lab: cx.cell(row=find_row(cx, lab, start=blk), column=6).value  # noqa: E731
    xl = dict(npv=X("NPV at discount rate"), irr=X("IRR"), pb=X("Payback period"),
              be=X("Break-even utilization of the added capacity"))
    d = [abs(xl["npv"] - py["npv"]), abs(xl["be"] - be), abs(npv_at_be)]
    if py["irr"] is not None:
        d.append(abs(xl["irr"] - py["irr"]))
    if py["pb"] is not None:
        d.append(abs(xl["pb"] - py["pb"]))
    else:
        d.append(0 if xl["pb"] == "Not within life" else 1)
    passes = py["npv"] > 0 and (py["irr"] or -1) >= P["hurdle"] and py["pb"] is not None and py["pb"] <= P["maxpb"]
    ok = max(d) < 1e-6
    ok_all &= ok
    out.append(dict(scenario=name, npv=round(py["npv"], 3), irr=None if py["irr"] is None else round(py["irr"], 4),
                    payback=None if py["pb"] is None else round(py["pb"], 3), break_even_util=round(be, 4),
                    npv_at_break_even=npv_at_be, meets_criteria=passes, max_abs_diff_vs_excel=max(d),
                    result="PASS" if ok else "FAIL"))
reco_xl = cx.cell(row=find_row(cx, "RECOMMENDATION"), column=3).value
down, base, up = (o["meets_criteria"] for o in out)
reco_py = ("REJECT" if not base else "APPROVE - meets" if down else "APPROVE WITH A DEMAND CHECKPOINT")
reco_ok = reco_xl.startswith(reco_py)
for o in out:
    print(o)
print("Recommendation (Excel):", reco_xl)
print("Recommendation rule re-derived in Python matches:", reco_ok)
(EN.tests / "phase5_capex_results.json").write_text(json.dumps(out, indent=2, default=str))
sys.exit(0 if ok_all and reco_ok else 1)

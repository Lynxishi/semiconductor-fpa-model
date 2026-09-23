"""
Phase 3 validation - independent Python rebuild of revenue, P&L and FCF for all three scenarios.

For each scenario: set Assumptions!C6, recalculate, rebuild everything in Python from the RAW inputs
(hardcoded cells only), compare every month of revenue, operating profit, net income, depreciation,
working capital and FCF, plus FY totals. Also confirms the Checks tab reads ALL CHECKS PASS.
Run: python tests/validate_phase3_scenarios.py
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_phase2_scenarios import PRODUCTS, find_row, read_inputs  # noqa: E402
from _common import EN, BY, FY, base_mix  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SRC = EN.model
RECALC = str(Path(__file__).resolve().parents[1] / "scripts" / "recalc.py")
TMP = EN.tests / "_scenario_runs3"
TMP.mkdir(exist_ok=True)


def model(a, g):
    idx = {"Downside": 0, "Base": 1, "Upside": 2}[a["scenario"]]
    d = {k: v[idx] for k, v in a["drivers"].items()}
    rev = [0.0] * 12
    cor = [0.0] * 12
    vcost = [0.0] * 12
    fcost = [0.0] * 12
    inv_close = [0.0] * 12
    base_vcost = inv0 = 0.0
    for p in PRODUCTS:
        x = a["prod"][p]
        rev0 = a["base_rev"] * x["mix"]
        sold0 = rev0 / x["asp0"]
        prod0 = sold0 / x["yld0"]
        base_vcost += prod0 * x["vc0"]
        cpgu0 = (prod0 * x["vc0"] + x["fix0"] * 12) / sold0
        dem_fy = sold0 * (1 + d["dem_g"] + x["gdiff"])
        asp, yld = x["asp0"] * (1 + d["asp_chg"]), min(1, max(0, x["yld0"] + d["yld_adj"]))
        vcu = x["vc0"] * (1 + d["mat_inf"])
        iu = x["inv_m"] * sold0 / 12
        iv = iu * cpgu0
        inv0 += iv
        for m, s in enumerate(a["season"]):
            produced = x["cap0"] * d["util"]
            good = produced * yld
            avail = iu + good
            sold = min(dem_fy * s / 12, avail)
            mfg = produced * vcu + x["fix0"]
            c = sold * (iv + mfg) / avail
            iu, iv = avail - sold, iv + mfg - c
            rev[m] += sold * asp
            cor[m] += c
            vcost[m] += produced * vcu
            fcost[m] += x["fix0"]
            inv_close[m] += iv
    capex = [rv * g["capex_pct"] for rv in rev]
    dep_new = [sum(capex[:m]) / g["life"] for m in range(12)]
    dep_fix = [f * g["dep_share"] for f in fcost]
    rd, sga = g["rd0"] * (1 + d["rd_g"]) / 12, g["sga0"] * (1 + d["sga_g"]) / 12
    op = [rev[m] - cor[m] - dep_new[m] - rd - sga for m in range(12)]
    ni = [o * (1 - g["tax"]) for o in op]
    ar = [rv * 12 * g["dso"] / 365 for rv in rev]
    ap = [v * 12 * g["dpo"] / 365 for v in vcost]
    wc = [ar[m] + inv_close[m] - ap[m] for m in range(12)]
    wc_open = a["base_rev"] * g["dso"] / 365 + inv0 - base_vcost * g["dpo"] / 365
    dwc = [wc[0] - wc_open] + [wc[m] - wc[m - 1] for m in range(1, 12)]
    da = [dep_new[m] + dep_fix[m] for m in range(12)]
    fcf = [ni[m] + da[m] - capex[m] - dwc[m] for m in range(12)]
    return dict(rev=rev, op=op, ni=ni, da=da, wc=wc, fcf=fcf)


def main():
    base = load_workbook(SRC, data_only=True)
    hd = base["Hist_Data"]
    mix = base_mix(hd)
    summary, all_ok = [], True
    for scen in ["Downside", "Base", "Upside"]:
        f = TMP / f"m_{scen}.xlsx"
        shutil.copy(SRC, f)
        wb = load_workbook(f)
        wb["Assumptions"]["C6"] = scen
        wb.save(f)
        res = json.loads(subprocess.run([sys.executable, RECALC, str(f), "120"], capture_output=True, text=True).stdout)
        v = load_workbook(f, data_only=True)
        asm = v["Assumptions"]
        a = read_inputs(asm, scen)
        sr = find_row(asm, "Seasonality factor (1.00 = average month)")
        a["season"] = [asm.cell(row=sr, column=c).value for c in range(3, 15)]
        for p in PRODUCTS:
            a["prod"][p]["mix"] = mix[p]
        idx = {"Downside": 3, "Base": 4, "Upside": 5}[scen]
        drv = lambda lab: asm.cell(row=find_row(asm, lab), column=idx).value  # noqa: E731
        fv = lambda lab: asm.cell(row=find_row(asm, lab), column=6).value     # noqa: E731
        g = {"capex_pct": drv("Capital expenditures (% of revenue, excl. proposed test-equipment project)"),
             "tax": drv("Tax rate"),
             "rd0": fv(f"Division R&D expense, {BY} base year"), "sga0": fv(f"Division SG&A expense, {BY} base year"),
             "dso": fv("Days sales outstanding (DSO)"), "dpo": fv("Days payable outstanding on variable manufacturing cost"),
             "dep_share": fv("Depreciation share of fixed manufacturing cost"), "life": fv(f"Useful life of new {FY} capex")}
        a["drivers"]["rd_g"] = [asm.cell(row=find_row(asm, "R&D expense growth"), column=c).value for c in (3, 4, 5)]
        a["drivers"]["sga_g"] = [asm.cell(row=find_row(asm, "SG&A expense growth"), column=c).value for c in (3, 4, 5)]
        py = model(a, g)

        pl, fc, rv, ck = v["PnL"], v["FCF"], v["Revenue"], v["Checks"]
        xl = {"rev": (pl, find_row(pl, "Revenue")), "op": (pl, find_row(pl, "Operating profit")),
              "ni": (pl, find_row(pl, "Net income")), "da": (fc, find_row(fc, "Total depreciation")),
              "wc": (fc, find_row(fc, "Operating working capital")), "fcf": (fc, find_row(fc, "Free cash flow"))}
        maxd, n = 0.0, 0
        for k, (ws, row) in xl.items():
            for m in range(12):
                maxd = max(maxd, abs(ws.cell(row=row, column=3 + m).value - py[k][m]))
                n += 1
        fy = {k: xl[k][0].cell(row=xl[k][1], column=15).value for k in ("rev", "op", "ni", "fcf")}
        fy_py = {"rev": sum(py["rev"]), "op": sum(py["op"]), "ni": sum(py["ni"]), "fcf": sum(py["fcf"])}
        fy_d = max(abs(fy[k] - fy_py[k]) for k in fy)
        overall = ck.cell(row=find_row(ck, "OVERALL MODEL STATUS"), column=3).value
        ok = maxd < 1e-6 and fy_d < 1e-6 and overall == "ALL CHECKS PASS" and res.get("total_errors") == 0
        all_ok &= ok
        summary.append(dict(scenario=scen, revenue=round(fy["rev"], 1), operating_profit=round(fy["op"], 1),
                            op_margin=round(fy["op"] / fy["rev"], 4), net_income=round(fy["ni"], 1),
                            fcf=round(fy["fcf"], 1), monthly_values_compared=n, max_abs_diff=maxd,
                            fy_max_abs_diff=fy_d, formula_errors=res.get("total_errors"), checks_tab=overall,
                            result="PASS" if ok else "FAIL"))
    for s in summary:
        print(s)
    (EN.tests / "phase3_scenario_results.json").write_text(json.dumps(summary, indent=2))
    shutil.rmtree(TMP, ignore_errors=True)  # Windows/OneDrive can briefly lock the temp copies
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()

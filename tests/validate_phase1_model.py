"""
Phase 1 validation: recompute the historical analysis in Python straight from the reported CSV
(independent of the workbook formulas) and compare with the recalculated Excel values.
Run after recalc: python tests/validate_phase1_model.py --company TICKER
"""
import csv
import sys

from openpyxl import load_workbook

from _common import EN

wb = load_workbook(EN.model, data_only=True)
ha, ck, hd = wb["Hist_Analysis"], wb["Checks"], wb["Hist_Data"]
rows = list(csv.DictReader(EN.reported_csv.open(encoding="utf-8")))
V = {(r["period"], r["line_id"]): float(r["value_usd_m"]) for r in rows if r["value_usd_m"] != ""}
YEARS = EN.years
COLS = "CDEFGHIJ"[:len(YEARS)]


def v(y, lid, default=None):
    if (y, lid) in V:
        return V[(y, lid)]
    if default is not None:
        return default
    raise KeyError((y, lid))


WC = sorted({lid for (_, lid) in V if lid.startswith("wc:")})
label_row = {ha.cell(row=r, column=1).value: r for r in range(1, ha.max_row + 1) if ha.cell(row=r, column=1).value}
seg0 = EN.products[0]["segment"]
seg0_label = next(hd.cell(row=r, column=1).value for r in range(1, hd.max_row + 1) if hd.cell(row=r, column=6 + len(YEARS)).value == seg0)

expected = {}
for i, y in enumerate(YEARS):
    ns, cogs = v(y, "net_sales"), v(y, "cost_of_sales")
    e = {}
    if i:
        e["Net sales growth"] = ns / v(YEARS[i - 1], "net_sales") - 1
    e["Gross margin"] = v(y, "gross_profit") / ns
    e["Operating margin (GAAP)"] = v(y, "operating_income") / ns
    e["Effective tax rate"] = v(y, "income_tax") / v(y, "pretax_income")
    wc = sum(v(y, k, 0.0) for k in WC)
    capex = -v(y, "cf:capex")
    e["Free cash flow (project definition)"] = v(y, "net_income") + v(y, "cf:da") - capex - (-wc)
    e["Free cash flow (CFO - capex, common market definition)"] = v(y, "cf:cfo") - capex
    e["Days inventory outstanding (DIO)"] = v(y, "bs:inventories") / cogs * 365
    e["Cash conversion cycle"] = (v(y, "bs:receivables") / ns + v(y, "bs:inventories") / cogs - v(y, "bs:payables") / cogs) * 365
    e[f"{seg0_label} - share of net sales"] = v(y, seg0) / ns
    e["Capex % of sales"] = capex / ns
    expected[y] = e

fails, n = 0, 0
print(f"{'Metric':58} {'Year':7} {'Excel':>12} {'Python':>12}  Result")
for y_i, y in enumerate(YEARS):
    for label, exp in expected[y].items():
        got = ha[f"{COLS[y_i]}{label_row[label]}"].value
        if got == "n.m.":
            ok = label == "Effective tax rate" and abs(v(y, "pretax_income")) < 0.02 * v(y, "net_sales")
            if ok:
                n += 1
                continue
        ok = isinstance(got, (int, float)) and abs(got - exp) < 1e-9
        n += 1
        fails += not ok
        if y in (YEARS[1], YEARS[-1]) or not ok:
            print(f"{label[:58]:58} {y:7} {got if isinstance(got, str) else f'{got:12.4f}':>12} {exp:12.4f}  {'PASS' if ok else 'FAIL'}")

statuses = [ck.cell(row=r, column=9).value for r in range(6, ck.max_row + 1)
            if ck.cell(row=r, column=9).value in ("PASS", "FAIL")]
overall = next(ck.cell(row=r, column=3).value for r in range(1, ck.max_row + 1)
               if ck.cell(row=r, column=1).value == "OVERALL MODEL STATUS")
print(f"\nPython vs Excel: {n - fails}/{n} match")
print(f"Checks tab: {statuses.count('PASS')} PASS, {statuses.count('FAIL')} FAIL -> {overall}")
sys.exit(1 if fails or statuses.count("FAIL") else 0)

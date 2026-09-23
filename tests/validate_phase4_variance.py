"""
Phase 4 validation - independent Python rebuild of budget, simulated actuals and the variance decomposition.

1. Reads only RAW inputs (hardcoded cells) from Assumptions and Actuals.
2. Rebuilds budget (Base scenario) and actuals month by month with its own code, then computes
   price / volume / mix and the operating-profit bridge.
3. Compares with the Excel Variance tab, with the selector on Base AND on Downside
   (budget and actuals must not move when the scenario changes).
Run: python tests/validate_phase4_variance.py
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_phase2_scenarios import PRODUCTS, PCOL, find_row, read_inputs  # noqa: E402
from _common import EN, BY, FY, base_mix  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SRC = EN.model
RECALC = str(Path(__file__).resolve().parents[1] / "scripts" / "recalc.py")
TMP = EN.tests / "_runs4"
TMP.mkdir(exist_ok=True)


def run_block(x, n, dem, util, yld, asp, vcu, fix, inv_u, inv_v):
    out = []
    for m in range(n):
        produced = x["cap0"] * util
        good = produced * yld
        avail = inv_u + good
        sold = min(dem[m], avail)
        mfg = produced * vcu + fix
        cor = sold * (inv_v + mfg) / avail
        inv_u, inv_v = avail - sold, inv_v + mfg - cor
        out.append(dict(dem=dem[m], prod=produced, good=good, sold=sold, rev=sold * asp, vcost=produced * vcu,
                        fix=fix, cor=cor, invv=inv_v))
    return out


def model(a, act, season, g):
    d = {k: v[1] for k, v in a["drivers"].items()}  # Base column
    bud, acu = {}, {}
    for p in PRODUCTS:
        x = a["prod"][p]
        sold0 = a["base_rev"] * x["mix"] / x["asp0"]
        cpgu0 = (sold0 / x["yld0"] * x["vc0"] + x["fix0"] * 12) / sold0
        inv_u = x["inv_m"] * sold0 / 12
        dem_b = [sold0 * (1 + d["dem_g"] + x["gdiff"]) * s / 12 for s in season]
        bud[p] = run_block(x, 12, dem_b, d["util"], min(1, max(0, x["yld0"] + d["yld_adj"])),
                           x["asp0"] * (1 + d["asp_chg"]), x["vc0"] * (1 + d["mat_inf"]), x["fix0"], inv_u, inv_u * cpgu0)
        q = act[p]
        dem_a = [dem_b[m] * (1 + q["a_dem"]) * act["noise"][m] for m in range(5)]
        acu[p] = run_block(x, 5, dem_a, q["a_util"], q["a_yld"], x["asp0"] * (1 + q["a_asp"]),
                           x["vc0"] * (1 + q["a_vc"]), q["a_fix"], inv_u, inv_u * cpgu0)
    rd_b, sga_b = g["rd0"] * (1 + d["rd_g"]) / 12, g["sga0"] * (1 + d["sga_g"]) / 12

    def div(blocks, n, capex_pct, rd, sga):
        rev = [sum(blocks[p][m]["rev"] for p in PRODUCTS) for m in range(n)]
        cor = [sum(blocks[p][m]["cor"] for p in PRODUCTS) for m in range(n)]
        capex = [x * capex_pct for x in rev]
        dep = [sum(capex[:m]) / g["life"] for m in range(n)]
        op = [rev[m] - cor[m] - dep[m] - rd - sga for m in range(n)]
        return dict(rev=rev, op=op, dep=dep)

    B = div(bud, 12, d["capex_pct"], rd_b, sga_b)
    A = div(acu, 5, act["a_capex"], act["a_rd"], act["a_sga"])
    Y = lambda blk, k, p: sum(blk[p][m][k] for m in range(5))  # noqa: E731
    Ub = {p: Y(bud, "sold", p) for p in PRODUCTS}
    Ua = {p: Y(acu, "sold", p) for p in PRODUCTS}
    TUb, TUa = sum(Ub.values()), sum(Ua.values())
    ASPb = {p: Y(bud, "rev", p) / Ub[p] for p in PRODUCTS}
    ASPa = {p: Y(acu, "rev", p) / Ua[p] for p in PRODUCTS}
    cb = {p: Y(bud, "cor", p) / Ub[p] for p in PRODUCTS}
    revb = sum(Y(bud, "rev", p) for p in PRODUCTS)
    res = {"price": sum(Ua[p] * (ASPa[p] - ASPb[p]) for p in PRODUCTS),
           "vol": (TUa - TUb) * revb / TUb,
           "mix": sum((Ua[p] / TUa - Ub[p] / TUb) * TUa * ASPb[p] for p in PRODUCTS)}
    res["rev_total"] = sum(A["rev"]) - sum(B["rev"][:5])
    yb = {p: Y(bud, "good", p) / Y(bud, "prod", p) for p in PRODUCTS}
    vcb = {p: Y(bud, "vcost", p) / Y(bud, "prod", p) for p in PRODUCTS}
    vca = {p: Y(acu, "vcost", p) / Y(acu, "prod", p) for p in PRODUCTS}
    res["mat"] = sum(-Y(acu, "prod", p) * (vca[p] - vcb[p]) for p in PRODUCTS)
    res["yld"] = sum(-vcb[p] * (Y(acu, "prod", p) - Y(acu, "good", p) / yb[p]) for p in PRODUCTS)
    res["opex"] = -((act["a_rd"] + act["a_sga"]) * 5 - (rd_b + sga_b) * 5)
    res["op_b"], res["op_a"] = sum(B["op"][:5]), sum(A["op"])
    res["op_var"] = res["op_a"] - res["op_b"]
    # full bridge rebuilt independently
    gpub = {p: ASPb[p] - cb[p] for p in PRODUCTS}
    vol_p = (TUa - TUb) * sum(Ub[p] / TUb * gpub[p] for p in PRODUCTS)
    mix_p = sum((Ua[p] / TUa - Ub[p] / TUb) * TUa * gpub[p] for p in PRODUCTS)
    fixv = -sum(Y(acu, "fix", p) - Y(bud, "fix", p) for p in PRODUCTS) - (sum(A["dep"]) - sum(B["dep"][:5]))
    absv = -sum(vcb[p] * (Y(acu, "good", p) / yb[p] - Y(bud, "prod", p)) - (acu[p][4]["invv"] - bud[p][4]["invv"])
                - (Ua[p] - Ub[p]) * cb[p] for p in PRODUCTS)
    res["bridge_end"] = res["op_b"] + res["price"] + vol_p + mix_p + res["mat"] + res["yld"] + fixv + absv + res["opex"]
    return res


def read_actual_inputs(ws):
    act = {p: {} for p in PRODUCTS}
    for key, lab in [("a_dem", "Demand vs budget"), ("a_asp", f"ASP change vs {BY} base-year ASP"),
                     ("a_yld", "Actual final-test yield"), ("a_util", "Actual capacity utilization"),
                     ("a_vc", f"Variable cost change vs {BY} base-year cost"),
                     ("a_fix", "Actual fixed manufacturing cost per month")]:
        rr = find_row(ws, lab)
        for p in PRODUCTS:
            act[p][key] = ws.cell(row=rr, column=PCOL[p]).value
    for key, lab in [("a_rd", "Actual R&D expense per month"), ("a_sga", "Actual SG&A expense per month"),
                     ("a_capex", "Actual capex % of revenue")]:
        act[key] = ws.cell(row=find_row(ws, lab), column=3).value
    rr = find_row(ws, "Monthly demand noise (applies to all products)")
    act["noise"] = [ws.cell(row=rr, column=c).value for c in range(3, 8)]
    return act


def main():
    summary, all_ok = [], True
    for scen in ["Base", "Downside"]:
        f = TMP / f"m_{scen}.xlsx"
        shutil.copy(SRC, f)
        wb = load_workbook(f)
        wb["Assumptions"]["C6"] = scen
        wb.save(f)
        rc = json.loads(subprocess.run([sys.executable, RECALC, str(f), "150"], capture_output=True, text=True).stdout)
        v = load_workbook(f, data_only=True)
        asm, va, ck, hd = v["Assumptions"], v["Variance"], v["Checks"], v["Hist_Data"]
        a = read_inputs(asm, scen)
        for k, lab in [("rd_g", "R&D expense growth"), ("sga_g", "SG&A expense growth"),
                       ("capex_pct", "Capital expenditures (% of revenue, excl. proposed test-equipment project)")]:
            a["drivers"][k] = [asm.cell(row=find_row(asm, lab), column=c).value for c in (3, 4, 5)]
        for p, m in base_mix(hd).items():
            a["prod"][p]["mix"] = m
        season = [asm.cell(row=find_row(asm, "Seasonality factor (1.00 = average month)"), column=c).value for c in range(3, 15)]
        fv = lambda lab: asm.cell(row=find_row(asm, lab), column=6).value  # noqa: E731
        g = {"rd0": fv(f"Division R&D expense, {BY} base year"), "sga0": fv(f"Division SG&A expense, {BY} base year"),
             "life": fv(f"Useful life of new {FY} capex")}
        py = model(a, read_actual_inputs(v["Actuals"]), season, g)
        X = lambda lab, col=6: va.cell(row=find_row(va, lab), column=col).value  # noqa: E731
        xl = {"price": X("Price variance"), "vol": X("Volume variance"), "mix": X("Product-mix variance"),
              "rev_total": X("Total revenue variance (actual - budget, computed directly)"),
              "mat": X("Material-cost variance"), "yld": X("Manufacturing-yield variance"),
              "opex": X("Operating-expense variance"), "op_var": X("Operating-profit variance"),
              "op_b": X("Budget operating profit (YTD)"), "op_a": X("Actual operating profit - from Actuals tab"),
              "bridge_end": X("Actual operating profit - from bridge")}
        diffs = {k: abs(xl[k] - py[k]) for k in xl}
        overall = ck.cell(row=find_row(ck, "OVERALL MODEL STATUS"), column=3).value
        ok = max(diffs.values()) < 1e-6 and overall == "ALL CHECKS PASS" and rc.get("total_errors") == 0
        all_ok &= ok
        summary.append(dict(selector=scen, **{k: round(xl[k], 3) for k in xl}, max_abs_diff_vs_python=max(diffs.values()),
                            formula_errors=rc.get("total_errors"), checks_tab=overall, result="PASS" if ok else "FAIL"))
    same = all(summary[0][k] == summary[1][k] for k in ("price", "vol", "mix", "op_b", "op_a", "op_var"))
    all_ok &= same
    for s in summary:
        print(s)
    print("Budget/actual/variance identical with selector on Base and Downside:", same)
    (EN.tests / "phase4_variance_results.json").write_text(json.dumps(summary, indent=2))
    shutil.rmtree(TMP, ignore_errors=True)  # Windows/OneDrive can briefly lock the temp copies
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()

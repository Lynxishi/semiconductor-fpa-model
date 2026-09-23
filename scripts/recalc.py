"""
Recalculate every formula in an .xlsx workbook with LibreOffice (headless) and report formula errors.

openpyxl writes formulas without cached results; LibreOffice opens the file, calculates it, and saves it back
as .xlsx with values. The file is replaced in place.

Usage:  python scripts/recalc.py model/Microchip_FPA_Model.xlsx [timeout_seconds]
Prints JSON: {"status": "success" | "errors_found", "total_formulas": n, "total_errors": n, "error_cells": [...]}
LibreOffice is found via the SOFFICE environment variable, the PATH, or the default Windows / macOS install folders.
(Alternative on Windows: open the workbook in Excel and save it - Excel recalculates on open.)
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
CANDIDATES = [r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"]


def find_soffice():
    for c in [os.environ.get("SOFFICE"), shutil.which("soffice"), shutil.which("libreoffice"), *CANDIDATES]:
        if c and Path(c).exists():
            return c
    return None


def recalc(path, timeout=180):
    path = Path(path).resolve()
    soffice = find_soffice()
    if not soffice:
        return {"error": "LibreOffice (soffice) not found. Install it or set SOFFICE, or open and save the file in Excel."}
    wb = load_workbook(path)
    n_formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("="))
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        cmd = [soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--calc",
               "--convert-to", "xlsx", "--outdir", tmp, str(path)]
        try:
            subprocess.run(cmd, capture_output=True, timeout=timeout, check=False)
        except subprocess.TimeoutExpired:
            return {"error": f"LibreOffice timed out after {timeout}s"}
        out = Path(tmp) / path.name
        if not out.exists():
            return {"error": "LibreOffice did not produce an output file"}
        shutil.copyfile(out, path)
    vals = load_workbook(path, data_only=True)
    bad = [f"{ws.title}!{c.coordinate}" for ws in vals.worksheets for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and c.value in ERRORS]
    return {"status": "success" if not bad else "errors_found", "total_formulas": n_formulas,
            "total_errors": len(bad), "error_cells": bad[:100]}


if __name__ == "__main__":
    res = recalc(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 180)
    print(json.dumps(res, indent=2))
    sys.exit(1 if "error" in res else 0)

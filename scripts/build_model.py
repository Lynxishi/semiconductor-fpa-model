"""
Builds outputs/<TICKER>/model/<TICKER>_FPA_Model.xlsx for the company chosen with --company (default MCHP).

Tabs: Cover, Sources_Log, Assumptions, Hist_Data, Hist_Analysis, Ops_Model, Revenue, PnL, FCF, Budget, Actuals,
Variance, Capex, Dash_Export, Checks. Everything company-specific comes from companies/<TICKER>/ (config.toml, the
reported CSV with standard line_ids, source register and assumptions log). Reported values are blue inputs with a
source on every row; every analysis figure is an Excel formula.

Run: python scripts/build_model.py --company MCHP   then   python scripts/recalc.py <model path>
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import engine  # noqa: E402

EN = engine.get()
ROOT = engine.ROOT
CFG = EN.cfg
OUT = EN.model
DISCLAIMER = EN.disclaimer
COMPANY = EN.company["name"]
YEARS = EN.years
YCOL = {y: get_column_letter(3 + i) for i, y in enumerate(YEARS)}  # C..(C+n-1)
FIRST, LAST = YCOL[YEARS[0]], YCOL[YEARS[-1]]
HLAB = CFG.get("history", {}).get("labels", {})
HNOTE = CFG.get("history", {}).get("notes", {})
SRCID = CFG.get("history", {}).get("source_ids", {})

# ---------- styles ----------
F = "Arial"
NAVY = "1F2A44"
f_base = Font(name=F, size=10)
f_input = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_calc = Font(name=F, size=10, color="000000")
f_bold = Font(name=F, size=10, bold=True)
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=16, bold=True, color=NAVY)
f_sub = Font(name=F, size=10, italic=True, color="595959")
f_sec = Font(name=F, size=10, bold=True, color=NAVY)
fill_hdr = PatternFill("solid", fgColor=NAVY)
fill_sec = PatternFill("solid", fgColor="E8ECF4")
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_pass = PatternFill("solid", fgColor="E2F0D9")
thin = Side(style="thin", color="BFBFBF")
b_top = Border(top=thin)

USD = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
DAYS = '0.0;(0.0);"-"'
MULT = '0.00x;(0.00x);"-"'

# ---------- load reported data (standard long format with line_id) ----------
rows = list(csv.DictReader(EN.reported_csv.open(encoding="utf-8")))
REP = {(r["period"], r["line_id"]): r for r in rows}
LIDS = []  # line ids in first-appearance order
for r_ in rows:
    if r_["line_id"] not in LIDS:
        LIDS.append(r_["line_id"])
LABEL = {}
for r_ in rows:
    LABEL.setdefault(r_["line_id"], r_["line_item"])
has = lambda lid: any((y, lid) in REP for y in YEARS)  # noqa: E731
comps = lambda prefix: [l for l in LIDS if l.startswith(prefix) and has(l)]  # noqa: E731
SEGMENTS = comps("segment:")
REQUIRED = ["net_sales", "cost_of_sales", "gross_profit", "rd", "sga", "operating_income", "pretax_income",
            "income_tax", "net_income", "cf:net_income", "cf:da", "cf:cfo", "cf:capex", "cf:cash_end", "bs:cash",
            "bs:receivables", "bs:inventories", "bs:total_current_assets", "bs:payables", "bs:total_current_liabilities"]
_missing = [l for l in REQUIRED if not has(l)]
if _missing:
    raise SystemExit(f"Reported CSV {EN.reported_csv.name} is missing required line_ids: {_missing}")
if not SEGMENTS:
    raise SystemExit("Reported CSV has no segment:* lines (needed for the product mix)")
for p_ in EN.products:
    if p_["segment"] not in SEGMENTS:
        raise SystemExit(f"Product {p_['key']}: segment {p_['segment']} not found in reported CSV (have {SEGMENTS})")


def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def header_row(ws, r, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=r, column=start_col + i, value=lab)
        c.font, c.fill = f_hdr, fill_hdr
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)


def section(ws, r, text, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=r, column=col).fill = fill_sec
    ws.cell(row=r, column=1, value=text).font = f_sec


def title_block(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = f_title
    ws["A2"] = subtitle
    ws["A2"].font = f_sub
    ws["A3"] = DISCLAIMER + "  (This tab contains reported public data and calculated values only.)" \
        if ws.title in ("Hist_Data", "Hist_Analysis") else DISCLAIMER
    ws["A3"].font = Font(name=F, size=9, italic=True, color="C00000")


wb = Workbook()

# =====================================================================
# Hist_Data - one row per reported line_id, grouped by statement
# =====================================================================
hd = wb.active
hd.title = "Hist_Data"
nY = len(YEARS)
C_CLS, C_SRC, C_NOTE, C_LID = 3 + nY, 4 + nY, 5 + nY, 6 + nY
title_block(hd, f"Historical Data - {COMPANY} (reported)",
            f"USD millions unless noted. Fiscal years end {EN.company['fiscal_year_end']}. Blue = hardcoded reported value from SEC filing.")
header_row(hd, 5, ["Line item", "Units"] + YEARS + ["Classification", "Source (filing, statement page)", "Notes", "line_id"])
setw(hd, {"A": 52, "B": 9, **{YCOL[y]: 11 for y in YEARS}, get_column_letter(C_CLS): 20, get_column_letter(C_SRC): 58,
          get_column_letter(C_NOTE): 60, get_column_letter(C_LID): 30})
hd.freeze_panes = "C6"
HD = {}  # line_id -> row

GROUPS = [
    ("Income statement", lambda l: l in ("net_sales", "cost_of_sales", "gross_profit", "rd", "sga", "total_opex",
                                         "operating_income", "pretax_income", "income_tax", "net_income")
     or l.startswith(("other_opex:", "nonop:", "memo:"))),
    ("Net sales by reported segment / product line", lambda l: l.startswith("segment:")),
    ("Other segment disclosures", lambda l: l.startswith("segment_info:")),
    ("Cash flow statement (selected lines)", lambda l: l.startswith(("cf:", "wc:"))),
    ("Balance sheet (period-end, selected lines)", lambda l: l.startswith("bs:")),
    ("Inventory detail", lambda l: l.startswith("inv:")),
]
IS_ORDER = ["net_sales", "cost_of_sales", "gross_profit", "rd", "sga"]
BOLD = ("gross_profit", "operating_income", "net_income", "cf:cfo", "bs:total_current_assets",
        "bs:total_current_liabilities", "bs:total_assets")


def ordered(pred):
    ids = [l for l in LIDS if pred(l) and has(l)]
    if pred("net_sales"):  # income statement: fixed skeleton, components where they belong
        body = [l for l in IS_ORDER if l in ids] + [l for l in ids if l.startswith("other_opex:")]
        body += [l for l in ("total_opex", "operating_income") if l in ids] + [l for l in ids if l.startswith("nonop:")]
        body += [l for l in ("pretax_income", "income_tax", "net_income") if l in ids] + [l for l in ids if l.startswith("memo:")]
        return body
    return ids


r = 6
for title, pred in GROUPS:
    ids = ordered(pred)
    if not ids:
        continue
    section(hd, r, title, C_LID)
    r += 1
    for lid in ids:
        HD[lid] = r
        hd.cell(row=r, column=1, value=HLAB.get(lid, LABEL[lid])).font = f_bold if lid in BOLD else f_base
        hd.cell(row=r, column=2, value="$m").font = f_base
        srcs = []
        for y in YEARS:
            rec = REP.get((y, lid))
            c = hd[f"{YCOL[y]}{r}"]
            if rec:
                c.value, c.font = float(rec["value_usd_m"]), f_input
                s = f"{SRCID.get(rec['accession_no'], rec['accession_no'])} {rec['source_url'].rsplit('/', 1)[-1]}"
                if s not in srcs:
                    srcs.append(s)
            c.number_format = USD
        hd.cell(row=r, column=C_CLS, value="Reported public data").font = f_base
        hd.cell(row=r, column=C_SRC, value="; ".join(srcs)).font = f_base
        hd.cell(row=r, column=C_NOTE, value=HNOTE.get(lid, "")).font = f_base
        hd.cell(row=r, column=C_LID, value=lid).font = Font(name=F, size=9, color="808080")
        if lid in BOLD:
            for col in range(1, 3 + nY):
                hd.cell(row=r, column=col).border = b_top
        r += 1

# Quarterly context (optional)
QROW, QCFG = {}, EN.company.get("quarter")
if QCFG and (QCFG["latest"], "net_sales") in REP:
    r += 1
    section(hd, r, f"Quarterly context - {QCFG['label']}", C_LID)
    r += 1
    header_row(hd, r, ["Line item", "Units", QCFG["prior"], QCFG["latest"]] + [""] * (nY - 2) + ["Classification", "Source", "Notes", "line_id"])
    r += 1
    for lid in ["net_sales"] + SEGMENTS:
        if (QCFG["latest"], lid) not in REP or (QCFG["prior"], lid) not in REP:
            continue
        QROW[lid] = r
        hd.cell(row=r, column=1, value=HLAB.get(lid, LABEL[lid])).font = f_base
        hd.cell(row=r, column=2, value="$m").font = f_base
        for col, p in [("C", QCFG["prior"]), ("D", QCFG["latest"])]:
            c = hd[f"{col}{r}"]
            c.value, c.font, c.number_format = float(REP[(p, lid)]["value_usd_m"]), f_input, USD
        rec = REP[(QCFG["latest"], lid)]
        hd.cell(row=r, column=C_CLS, value="Reported public data").font = f_base
        hd.cell(row=r, column=C_SRC, value=f"{SRCID.get(rec['accession_no'], rec['accession_no'])} {rec['source_url'].rsplit('/', 1)[-1]}").font = f_base
        hd.cell(row=r, column=C_NOTE, value=f"{QCFG['prior']} = prior-year comparative column" if lid == "net_sales" else "").font = f_base
        hd.cell(row=r, column=C_LID, value=lid).font = Font(name=F, size=9, color="808080")
        r += 1

# Operating facts (optional, percentages for the base year)
FACT_ROW = {}
facts = [l for l in LIDS if l.startswith("fact:") and (EN.base_year, l) in REP]
if facts:
    r += 1
    section(hd, r, f"Operating facts stated in the {EN.base_year} annual report (approximate, as stated)", C_LID)
    r += 1
    for lid in facts:
        FACT_ROW[lid] = r
        rec = REP[(EN.base_year, lid)]
        hd.cell(row=r, column=1, value=LABEL[lid].replace(" (%)", "")).font = f_base
        hd.cell(row=r, column=2, value="%").font = f_base
        c = hd[f"{YCOL[EN.base_year]}{r}"]
        c.value, c.font, c.number_format = float(rec["value_usd_m"]) / 100, f_input, '0%'
        hd.cell(row=r, column=C_CLS, value="Reported public data").font = f_base
        hd.cell(row=r, column=C_SRC, value=f"{SRCID.get(rec['accession_no'], rec['accession_no'])} (annual report, Item 1)").font = f_base
        hd.cell(row=r, column=C_NOTE, value=rec.get("note", "")).font = f_base
        hd.cell(row=r, column=C_LID, value=lid).font = Font(name=F, size=9, color="808080")
        r += 1


def H(lid, y):
    """Absolute reference to a Hist_Data cell (by line_id)."""
    return f"Hist_Data!${YCOL[y]}${HD[lid]}"


def HN(lid, y):
    """Reference that reads 0 when the line is not reported (N() of a blank)."""
    return f"N({H(lid, y)})" if lid in HD else "0"


def HSUM(prefix, y):
    ids = comps(prefix)
    return "(" + "+".join(HN(l, y) for l in ids) + ")" if ids else "0"


# =====================================================================
# Hist_Analysis
# =====================================================================
ha = wb.create_sheet("Hist_Analysis")
title_block(ha, f"Historical Financial Analysis - {COMPANY}",
            "All figures are formulas on Hist_Data. Green = direct link to Hist_Data; black = calculation.")
header_row(ha, 5, ["Metric", "Units"] + YEARS + ["Classification", "Formula / definition"])
C_HCLS, C_HDEF = 3 + nY, 4 + nY
setw(ha, {"A": 50, "B": 9, **{YCOL[y]: 11 for y in YEARS}, get_column_letter(C_HCLS): 18, get_column_letter(C_HDEF): 70})
ha.freeze_panes = "C6"
HA = {}
r = 6


def ha_row(key, label, units, fn, fmt, definition, first_blank=False, link=False, bold=False):
    """fn(y, prev_y) -> formula string (without '=')."""
    global r
    HA[key] = r
    ha.cell(row=r, column=1, value=label).font = f_bold if bold else f_base
    ha.cell(row=r, column=2, value=units).font = f_base
    for i, y in enumerate(YEARS):
        c = ha[f"{YCOL[y]}{r}"]
        if first_blank and i == 0:
            c.value = "n/a"
            c.font = Font(name=F, size=9, color="808080")
            c.alignment = Alignment(horizontal="right")
            continue
        prev = YEARS[i - 1] if i else None
        c.value = "=" + fn(y, prev)
        c.font = f_link if link else f_calc
        c.number_format = fmt
    ha.cell(row=r, column=C_HCLS, value="Calculated value").font = f_base
    ha.cell(row=r, column=C_HDEF, value=definition).font = f_base
    r += 1


def A(key, y):
    return f"${YCOL[y]}${HA[key]}"


def AA(key, y):
    """Absolute reference to a Hist_Analysis cell, for use from other tabs."""
    return f"Hist_Analysis!{A(key, y)}"


def sec(text):
    global r
    section(ha, r, text, C_HDEF)
    r += 1


OO = comps("other_opex:")
OO_LABEL = ", ".join(HLAB.get(l, LABEL[l]).lower() for l in OO) or "none reported"
sec("Growth and profitability")
ha_row("sales", "Net sales", "$m", lambda y, p: H("net_sales", y), USD, "Link: Hist_Data net sales", link=True)
ha_row("g_sales", "Net sales growth", "%", lambda y, p: f"IFERROR({H('net_sales', y)}/{H('net_sales', p)}-1,\"\")",
       PCT, "Net sales / prior-year net sales - 1", first_blank=True)
ha_row("gm", "Gross margin", "%", lambda y, p: f"IFERROR({H('gross_profit', y)}/{H('net_sales', y)},\"\")", PCT,
       "Gross profit / net sales")
ha_row("rd_pct", "R&D % of sales", "%", lambda y, p: f"IFERROR({H('rd', y)}/{H('net_sales', y)},\"\")", PCT,
       "Research and development / net sales")
ha_row("sga_pct", "SG&A % of sales", "%", lambda y, p: f"IFERROR({H('sga', y)}/{H('net_sales', y)},\"\")", PCT,
       "SG&A / net sales")
ha_row("om", "Operating margin (GAAP)", "%", lambda y, p: f"IFERROR({H('operating_income', y)}/{H('net_sales', y)},\"\")", PCT,
       "Operating income / net sales")
ha_row("oo_total", "Other operating items (sum of reported components)", "$m", lambda y, p: HSUM("other_opex:", y), USD,
       f"Sum of: {OO_LABEL}")
ha_row("oi_core", "Operating income before other operating items", "$m",
       lambda y, p: f"{H('operating_income', y)}+{A('oo_total', y)}", USD,
       f"Operating income + other operating items ({OO_LABEL}). Analyst calculation - NOT {COMPANY}'s reported non-GAAP measure.")
ha_row("om_core", "Margin before other operating items", "%",
       lambda y, p: f"IFERROR({A('oi_core', y)}/{H('net_sales', y)},\"\")", PCT,
       "Row above / net sales. Closest proxy for a product division's operating margin.")
ha_row("nm", "Net margin", "%", lambda y, p: f"IFERROR({H('net_income', y)}/{H('net_sales', y)},\"\")", PCT,
       "Net income / net sales")
ha_row("etr", "Effective tax rate", "%",
       lambda y, p: f"IF(ABS({H('pretax_income', y)})<0.02*{H('net_sales', y)},\"n.m.\",IFERROR({H('income_tax', y)}/{H('pretax_income', y)},\"\"))",
       PCT, "Income tax / income before taxes. Shows n.m. (not meaningful) when pre-tax income is under 2% of sales.")
ha_row("incr_gm", "Incremental gross margin (change in GP / change in sales)", "%",
       lambda y, p: f"IFERROR(({H('gross_profit', y)}-{H('gross_profit', p)})/({H('net_sales', y)}-{H('net_sales', p)}),\"\")", PCT,
       "Shows operating leverage: how much gross profit moved per $1 of sales change (fixed-cost absorption)",
       first_blank=True)

sec("Reported segment / product-line mix and growth")
for lid in SEGMENTS:
    lab = HLAB.get(lid, LABEL[lid])
    ha_row(f"mix:{lid}", f"{lab} - share of net sales", "%",
           lambda y, p, lid=lid: f"IFERROR({H(lid, y)}/{H('net_sales', y)},\"\")", PCT, f"{lab} sales / net sales")
ha_row("mix_total", "Mix total (must equal 100%)", "%", lambda y, p: "+".join(A(f"mix:{l}", y) for l in SEGMENTS), PCT,
       "Sum of the segment shares", bold=True)
for lid in SEGMENTS:
    lab = HLAB.get(lid, LABEL[lid])
    ha_row(f"g:{lid}", f"{lab} - growth", "%",
           lambda y, p, lid=lid: f"IFERROR({H(lid, y)}/{H(lid, p)}-1,\"\")", PCT, f"{lab} sales / prior year - 1",
           first_blank=True)

WC = comps("wc:")
sec("Cash flow and free cash flow (project definition)")
ha_row("fcf_ni", "Net income", "$m", lambda y, p: H("net_income", y), USD, "Link: Hist_Data net income", link=True)
ha_row("fcf_da", "Depreciation and amortization", "$m", lambda y, p: H("cf:da", y), USD,
       "Link: Hist_Data D&A (cash flow statement)", link=True)
ha_row("capex_pos", "Capital expenditures (positive = spend)", "$m", lambda y, p: f"-{H('cf:capex', y)}", USD,
       "Negative of capex as printed. Cash-flow statement prints capex negative.")
ha_row("inc_wc", "Increase in working capital", "$m", lambda y, p: f"-{HSUM('wc:', y)}", USD,
       f"Negative of the sum of the {len(WC)} cash-flow working-capital lines. Positive = cash absorbed by working capital.")
ha_row("fcf", "Free cash flow (project definition)", "$m",
       lambda y, p: f"{A('fcf_ni', y)}+{A('fcf_da', y)}-{A('capex_pos', y)}-{A('inc_wc', y)}", USD,
       "Net income + D&A - capex - increase in working capital (project formula)", bold=True)
ha_row("noncash_oth", "Other non-cash items in CFO (SBC, deferred tax, other)", "$m",
       lambda y, p: f"{H('cf:cfo', y)}-{A('fcf_ni', y)}-{A('fcf_da', y)}+{A('inc_wc', y)}", USD,
       "CFO - net income - D&A + increase in WC. Explains the gap to the CFO-based FCF below.")
ha_row("fcf_cfo", "Free cash flow (CFO - capex, common market definition)", "$m",
       lambda y, p: f"{H('cf:cfo', y)}-{A('capex_pos', y)}", USD,
       "Net cash from operations - capex. Reconciles to project FCF + other non-cash items (see Checks).")
ha_row("fcf_m", "FCF margin (project definition)", "%",
       lambda y, p: f"IFERROR({A('fcf', y)}/{H('net_sales', y)},\"\")", PCT, "Project FCF / net sales")
ha_row("fcf_conv", "FCF conversion (project FCF / net income)", "x",
       lambda y, p: f"IF(ABS({A('fcf_ni', y)})<0.02*{H('net_sales', y)},\"n.m.\",IFERROR({A('fcf', y)}/{A('fcf_ni', y)},\"\"))", MULT,
       "Project FCF / net income. Shows n.m. when net income is under 2% of sales.")
ha_row("capex_pct", "Capex % of sales", "%", lambda y, p: f"IFERROR({A('capex_pos', y)}/{H('net_sales', y)},\"\")",
       PCT, "Capex / net sales - anchor for the forecast capex assumption")
ha_row("da_pct", "D&A % of sales", "%", lambda y, p: f"IFERROR({H('cf:da', y)}/{H('net_sales', y)},\"\")", PCT,
       "D&A / net sales")
if has("cf:sbc"):
    ha_row("sbc_pct", "Share-based compensation % of sales", "%",
           lambda y, p: f"IFERROR({H('cf:sbc', y)}/{H('net_sales', y)},\"\")", PCT, "SBC / net sales")
ha_row("capex_da", "Capex / depreciation & amortization", "x",
       lambda y, p: f"IFERROR({A('capex_pos', y)}/{H('cf:da', y)},\"\")", MULT,
       "<1.0x means spending less than D&A (asset base shrinking or amortization-heavy D&A)")

sec("Working capital and asset efficiency (year-end balances)")
ha_row("dso", "Days sales outstanding (DSO)", "days",
       lambda y, p: f"IFERROR({H('bs:receivables', y)}/{H('net_sales', y)}*365,\"\")", DAYS, "Accounts receivable / net sales x 365")
ha_row("dio", "Days inventory outstanding (DIO)", "days",
       lambda y, p: f"IFERROR({H('bs:inventories', y)}/{H('cost_of_sales', y)}*365,\"\")", DAYS,
       "Inventories / annual cost of sales x 365 (company-reported days often use quarterly COGS, so they can differ).")
ha_row("dpo", "Days payable outstanding (DPO)", "days",
       lambda y, p: f"IFERROR({H('bs:payables', y)}/{H('cost_of_sales', y)}*365,\"\")", DAYS, "Accounts payable / cost of sales x 365")
ha_row("ccc", "Cash conversion cycle", "days", lambda y, p: f"{A('dso', y)}+{A('dio', y)}-{A('dpo', y)}", DAYS,
       "DSO + DIO - DPO", bold=True)
ha_row("nwc", "Operating working capital", "$m",
       lambda y, p: f"{H('bs:receivables', y)}+{H('bs:inventories', y)}+{HN('bs:other_current_assets', y)}-{H('bs:payables', y)}-{HN('bs:accrued', y)}",
       USD, "AR + inventories + other current assets - AP - accrued liabilities (excludes cash and debt)")
ha_row("nwc_pct", "Operating working capital % of sales", "%",
       lambda y, p: f"IFERROR({A('nwc', y)}/{H('net_sales', y)},\"\")", PCT, "Anchor for the forecast working-capital assumption")
if has("bs:ppe"):
    ha_row("ppe_turn", "Net sales / net PP&E", "x", lambda y, p: f"IFERROR({H('net_sales', y)}/{H('bs:ppe', y)},\"\")", MULT,
           "Fixed-asset turnover")

if QROW:
    sec(f"Quarterly context ({QCFG['label']})")
    for i, lab in enumerate(["", "", QCFG["prior"], QCFG["latest"], "YoY growth"]):
        if lab:
            c = ha.cell(row=r, column=1 + i, value=lab)
            c.font = f_bold
            c.alignment = Alignment(horizontal="right")
    r += 1
    for lid in QROW:
        HA[f"q:{lid}"] = r
        ha.cell(row=r, column=1, value=HLAB.get(lid, LABEL[lid])).font = f_base
        ha.cell(row=r, column=2, value="$m").font = f_base
        for col in ["C", "D"]:
            c = ha[f"{col}{r}"]
            c.value = f"=Hist_Data!${col}${QROW[lid]}"
            c.font, c.number_format = f_link, USD
        c = ha[f"E{r}"]
        c.value = f"=IFERROR(D{r}/C{r}-1,\"\")"
        c.font, c.number_format = f_calc, PCT
        ha.cell(row=r, column=C_HCLS, value="Calculated value").font = f_base
        ha.cell(row=r, column=C_HDEF, value=f"{QCFG['latest']} / {QCFG['prior']} - 1. Context for the {EN.fy} demand-growth assumption.").font = f_base
        r += 1

# Analyst observations (text built by formulas; the years quoted are chosen from the data)
sales = {y: float(REP[(y, "net_sales")]["value_usd_m"]) for y in YEARS}
peak = max(YEARS, key=lambda y: sales[y])
after = YEARS[YEARS.index(peak):]
trough = min(after, key=lambda y: sales[y])
last, first = YEARS[-1], YEARS[0]
r += 1
section(ha, r, "Analyst observations (each refers to a metric above; numbers quoted are formula results)", C_HDEF)
r += 1
OBS = []
if trough != peak:
    OBS.append(("Cycle", f"=\"Net sales peaked in {peak} at $\"&TEXT({H('net_sales', peak)},\"#,##0\")&\"m and fell \"&TEXT(-({H('net_sales', trough)}/{H('net_sales', peak)}-1),\"0%\")&\" by {trough}\"" +
                (f"&\", then grew \"&TEXT({A('g_sales', last)},\"0.0%\")&\" in {last}.\"" if trough != last else "&\".\"")))
    OBS.append(("Operating leverage", f"=\"Gross margin moved from \"&TEXT({A('gm', peak)},\"0.0%\")&\" ({peak}) to \"&TEXT({A('gm', trough)},\"0.0%\")&\" ({trough}): consistent with fixed factory cost spread over fewer units. The operating model simulates this effect.\""))
else:
    OBS.append(("Growth", f"=\"Net sales grew from $\"&TEXT({H('net_sales', first)},\"#,##0\")&\"m in {first} to $\"&TEXT({H('net_sales', last)},\"#,##0\")&\"m in {last}.\""))
OBS.append(("Inventory", f"=\"Days of inventory were \"&TEXT({A('dio', first)},\"0\")&\" in {first} and \"&TEXT({A('dio', last)},\"0\")&\" in {last}.\""))
OBS.append(("Capex", f"=\"Capex was \"&TEXT({A('capex_pct', first)},\"0.0%\")&\" of sales in {first} and \"&TEXT({A('capex_pct', last)},\"0.0%\")&\" in {last}; capex / D&A was \"&TEXT({A('capex_da', last)},\"0.00x\")&\" in {last}.\""))
if QROW:
    OBS.append(("Latest quarter", f"=\"{QCFG['latest']} net sales changed \"&TEXT(E{HA['q:net_sales']},\"+0%;-0%\")&\" year over year.\""))
for lab, formula in OBS:
    ha.cell(row=r, column=1, value=lab).font = f_bold
    ha.cell(row=r, column=3, value=formula).font = f_calc
    r += 1

# =====================================================================
# Checks (Phase 1 checks; later phases append)
# =====================================================================
ck = wb.create_sheet("Checks")
title_block(ck, "Model Checks", "Each check computes an independent result and compares it to the model. "
                                 "Tolerance = 0.05 ($m) or 0.0001 (ratios).")
TOLC, STC = get_column_letter(3 + nY), get_column_letter(4 + nY)
header_row(ck, 5, ["Check", "Area"] + YEARS + ["Tolerance", "Status"])
setw(ck, {"A": 60, "B": 16, **{YCOL[y]: 11 for y in YEARS}, TOLC: 11, STC: 12})
ck.freeze_panes = "C6"
# The later phases write single-result checks with result in C, tolerance in H and status in I.
# Keep those fixed columns: for 5 history years they coincide with the year grid (C..G, H, I).
assert nY == 5, "history_years must list 5 fiscal years (the Checks layout assumes 5)"
r = 6
CK_ROWS = []


def ck_row(label, area, diff_fn, tol=0.05, years=YEARS):
    """diff_fn(y) -> formula for (model - independent). Status PASS if all |diff| <= tol."""
    global r
    ck.cell(row=r, column=1, value=label).font = f_base
    ck.cell(row=r, column=2, value=area).font = f_base
    for y in YEARS:
        c = ck[f"{YCOL[y]}{r}"]
        if y in years:
            c.value = "=" + diff_fn(y)
            c.number_format = '0.0000;(0.0000);"-"'
            c.font = f_calc
        else:
            c.value = "n/a"
            c.font = Font(name=F, size=9, color="808080")
            c.alignment = Alignment(horizontal="right")
    t = ck[f"H{r}"]
    t.value, t.font = tol, f_input
    cond = ",".join(f"ABS({YCOL[y]}{r})<=$H${r}" for y in years)
    s = ck[f"I{r}"]
    s.value = f'=IF(AND({cond}),"PASS","FAIL")'
    s.font = f_bold
    CK_ROWS.append(r)
    r += 1


def cksec(text):
    global r
    section(ck, r, text, 9)
    r += 1


cksec("Reported data - internal consistency")
ck_row("Gross profit = net sales - cost of sales", "Income statement",
       lambda y: f"{H('gross_profit', y)}-({H('net_sales', y)}-{H('cost_of_sales', y)})")
if has("total_opex"):
    ck_row("Operating expenses = R&D + SG&A + other operating items", "Income statement",
           lambda y: f"{H('total_opex', y)}-({H('rd', y)}+{H('sga', y)}+{HSUM('other_opex:', y)})")
    ck_row("Operating income = gross profit - operating expenses", "Income statement",
           lambda y: f"{H('operating_income', y)}-({H('gross_profit', y)}-{H('total_opex', y)})")
else:  # no printed total: test the components directly
    ck_row("Operating income = gross profit - R&D - SG&A - other operating items", "Income statement",
           lambda y: f"{H('operating_income', y)}-({H('gross_profit', y)}-{H('rd', y)}-{H('sga', y)}-{HSUM('other_opex:', y)})")
ck_row("Pre-tax income = operating income + non-operating items", "Income statement",
       lambda y: f"{H('pretax_income', y)}-({H('operating_income', y)}+{HSUM('nonop:', y)})")
ck_row("Net income = pre-tax income - tax", "Income statement",
       lambda y: f"{H('net_income', y)}-({H('pretax_income', y)}-{H('income_tax', y)})")
ck_row("Reported segments / product lines sum to net sales", "Revenue",
       lambda y: f"{HSUM('segment:', y)}-{H('net_sales', y)}")
seg_tot = [l for l in EN.company.get("checks", {}).get("segments_sum_to_sales", []) if has(l)]
if seg_tot:
    yrs = [y for y in YEARS if all((y, l) in REP for l in seg_tot)]
    ck_row("Reportable segments sum to net sales", "Revenue",
           lambda y: "(" + "+".join(H(l, y) for l in seg_tot) + f")-{H('net_sales', y)}", years=yrs)
ck_row("Cash-flow net income = income-statement net income", "Cash flow",
       lambda y: f"{H('cf:net_income', y)}-{H('net_income', y)}")
ck_row("Cash-flow ending cash = balance-sheet cash", "Cash flow", lambda y: f"{H('cf:cash_end', y)}-{H('bs:cash', y)}")
CA_PARTS = ["bs:cash", "bs:st_investments", "bs:receivables", "bs:inventories", "bs:other_current_assets"]
CL_PARTS = ["bs:payables", "bs:accrued", "bs:current_debt"]
ck_row("Total current assets = cash + ST investments + AR + inventory + other (+ other listed)", "Balance sheet",
       lambda y: f"{H('bs:total_current_assets', y)}-(" + "+".join(HN(l, y) for l in CA_PARTS) + f"+{HSUM('bs:ca_other:', y)})")
ck_row("Total current liabilities = AP + accrued + current debt (+ other listed)", "Balance sheet",
       lambda y: f"{H('bs:total_current_liabilities', y)}-(" + "+".join(HN(l, y) for l in CL_PARTS) + f"+{HSUM('bs:cl_other:', y)})")
if comps("inv:"):
    yrs = [y for y in YEARS if all((y, l) in REP for l in comps("inv:"))]
    ck_row("Inventory components sum to total inventory", "Balance sheet",
           lambda y: f"{H('bs:inventories', y)}-{HSUM('inv:', y)}", years=yrs)

cksec("Historical analysis - independent recomputation")
ck_row("Gross margin = 1 - cost of sales / net sales (alternate route)", "Analysis",
       lambda y: f"{AA('gm', y)}-(1-{H('cost_of_sales', y)}/{H('net_sales', y)})", tol=0.0001)
ck_row("Segment mix sums to 100%", "Analysis", lambda y: f"{AA('mix_total', y)}-1", tol=0.0001)
ck_row("Growth chain: prior-year sales x (1 + growth) = current sales", "Analysis",
       lambda y: f"{H('net_sales', YEARS[YEARS.index(y) - 1])}*(1+{AA('g_sales', y)})-{H('net_sales', y)}",
       years=YEARS[1:])
ck_row("FCF (project) recomputed directly from Hist_Data", "Free cash flow",
       lambda y: f"{AA('fcf', y)}-({H('net_income', y)}+{H('cf:da', y)}+{H('cf:capex', y)}+{HSUM('wc:', y)})")
ck_row("FCF (project) + other non-cash items = CFO - capex", "Free cash flow",
       lambda y: f"({AA('fcf', y)}+{AA('noncash_oth', y)})-({H('cf:cfo', y)}+{H('cf:capex', y)})")
ck_row("Cash conversion cycle = DSO + DIO - DPO (recomputed from balances)", "Working capital",
       lambda y: f"{AA('ccc', y)}-(({H('bs:receivables', y)}/{H('net_sales', y)}+{H('bs:inventories', y)}/{H('cost_of_sales', y)}-{H('bs:payables', y)}/{H('cost_of_sales', y)})*365)",
       tol=0.0001)

# ---- Phase 2+ tabs (Assumptions ... Dash_Export), built here so the Hist helpers (H, AA) exist
_r_ck = r
for mod, fn in [("phase2_ops", "build_phase2_tabs"), ("phase3_pnl", "build_phase3_tabs"),
                ("phase4_variance", "build_phase4_tabs"), ("phase5_capex", "build_phase5_tab"),
                ("phase6_export", "build_phase6_tab")]:
    exec(open(ROOT / "scripts" / f"{mod}.py", encoding="utf-8").read())
    globals()[fn]()
r = _r_ck

cksec("Hand-calculable unit tests (fixed inputs, known answers)")
tests = [
    ("Growth: 100 -> 110 = 10%", "=110/100-1", 0.10),
    ("Gross margin: sales 200, COGS 80 = 60%", "=(200-80)/200", 0.60),
    ("FCF: NI 100 + D&A 20 - capex 30 - incr. WC 10 = 80", "=100+20-30-10", 80),
    ("Increase in WC from CF lines: AR -15, inventory +5, AP +3 -> +7 absorbed", "=-(-15+5+3)", 7),
    ("DSO: AR 100, sales 365 = 100 days", "=100/365*365", 100),
    ("CCC: DSO 50 + DIO 120 - DPO 30 = 140 days", "=50+120-30", 140),
    ("Mix: 50 + 30 + 20 of 100 sums to 100%", "=50/100+30/100+20/100", 1),
]
header_row(ck, r, ["Unit test", "", "Result", "Expected", "Difference", "", "", "Tolerance", "Status"])
r += 1
for lab, formula, exp in tests:
    ck.cell(row=r, column=1, value=lab).font = f_base
    ck[f"C{r}"] = formula
    ck[f"D{r}"] = exp
    ck[f"D{r}"].font = f_input
    ck[f"E{r}"] = f"=C{r}-D{r}"
    ck[f"H{r}"] = 0.000001
    ck[f"H{r}"].font = f_input
    ck[f"I{r}"] = f'=IF(ABS(E{r})<=H{r},"PASS","FAIL")'
    ck[f"I{r}"].font = f_bold
    CK_ROWS.append(r)
    r += 1

add_phase2_checks()  # noqa: F821 (defined by the exec'd phase modules)
add_phase3_checks()  # noqa: F821
add_phase4_checks()  # noqa: F821
add_phase5_checks()  # noqa: F821
add_phase6_checks()  # noqa: F821

r += 1
ck.cell(row=r, column=1, value="OVERALL MODEL STATUS").font = Font(name=F, size=11, bold=True)
ck[f"C{r}"] = f'=IF(COUNTIF(I6:I{r - 1},"FAIL")=0,"ALL CHECKS PASS","CHECK FAILURES: "&COUNTIF(I6:I{r - 1},"FAIL"))'
ck[f"C{r}"].font = Font(name=F, size=11, bold=True)
ck[f"C{r}"].fill = fill_key
CK_STATUS = f"Checks!$C${r}"
_meta_status = next(rr for rr in range(DX["meta"]["first"], DX["meta"]["last"] + 1)  # noqa: F821
                    if wb["Dash_Export"].cell(row=rr, column=1).value == "checks_status")
wb["Dash_Export"].cell(row=_meta_status, column=2, value=f"={CK_STATUS}").font = f_link
ck[f"G{r}"] = f'=COUNTIF(I6:I{r - 1},"PASS")&" pass / "&COUNTIF(I6:I{r - 1},"FAIL")&" fail"'
ck[f"G{r}"].font = f_bold

# =====================================================================
# Sources_Log
# =====================================================================
sl = wb.create_sheet("Sources_Log", 0)
title_block(sl, "Source Register and Assumptions Log", "Every reported value traces to a source below. "
                                                         "Every assumption is logged with who decided it.")
setw(sl, {"A": 10, "B": 30, "C": 26, "D": 50, "E": 60, "F": 22, "G": 22, "H": 12})
r = 5
section(sl, r, "Source register", 8)
r += 1
header_row(sl, r, ["ID", "Document", "Accession no.", "Period covered", "URL", "Used for", "Accessed", ""])
r += 1
for x in csv.DictReader(EN.source_register.open(encoding="utf-8")):
    vals = [x["source_id"], x["document"], x["accession_no"], x["period_covered"], x["url"], x["used_for"], x["accessed"]]
    for i, v in enumerate(vals):
        c = sl.cell(row=r, column=1 + i, value=v)
        c.font, c.alignment = f_base, Alignment(wrap_text=True, vertical="top")
    sl.cell(row=r, column=5).hyperlink = x["url"]
    r += 1
r += 1
section(sl, r, "Assumptions log", 8)
r += 1
header_row(sl, r, ["ID", "Area", "Assumption", "Value", "Rationale", "Classification", "Decided by", "Status"])
r += 1
for x in csv.DictReader(EN.assumptions_log.open(encoding="utf-8")):
    vals = [x["id"], x["area"], x["assumption"], x["value"], x["rationale"], x["classification"], x["decided_by"], x["status"]]
    for i, v in enumerate(vals):
        c = sl.cell(row=r, column=1 + i, value=v)
        c.font, c.alignment = f_base, Alignment(wrap_text=True, vertical="top")
    r += 1

# =====================================================================
# Cover
# =====================================================================
co, dv, tp, fyc = EN.company, EN.division, EN.target, CFG["forecast"]
cv = wb.create_sheet("Cover", 0)
setw(cv, {"A": 3, "B": 28, "C": 90})
cv["B2"] = "Semiconductor FP&A Forecasting and Decision Model"
cv["B2"].font = Font(name=F, size=20, bold=True, color=NAVY)
cv["B3"] = f"Fictional {dv['name']}, modeled on {co['legal_name']} ({co['exchange']}: {co['ticker']})"
cv["B3"].font = Font(name=F, size=11, color="595959")
cv["B5"] = DISCLAIMER
cv["B5"].font = Font(name=F, size=10, bold=True, color="C00000")
cv["B6"] = (f"Reported public data comes only from {COMPANY}'s SEC filings (see Sources_Log). "
            "All division-level operating data is simulated and labeled as such.")
cv["B6"].font = f_sub
cv["B8"] = "Business question"
cv["B8"].font = f_sec
cv["C8"] = ("How would changes in demand, average selling price, product mix, manufacturing yield, capacity "
            "utilization, operating costs and capital spending affect the division's revenue, operating profit "
            "and free cash flow?")
cv["C8"].alignment = Alignment(wrap_text=True, vertical="top")
cv["C8"].font = f_base
cv.row_dimensions[8].height = 42
cv["B10"] = "Model status"
cv["B10"].font = f_sec
cv["C10"] = f"={CK_STATUS}"
cv["C10"].font = Font(name=F, size=10, bold=True)
cv["B11"] = "Built by"
cv["B11"].font = f_sec
cv["C11"] = f"scripts/build_model.py --company {EN.ticker} (config: companies/{EN.ticker}/config.toml)"
cv["C11"].font = f_base
cv["B13"] = "Tab guide"
cv["B13"].font = f_sec
tabs = [("Cover", "This page"),
        ("Sources_Log", "Source register (SEC filings) and assumptions log"),
        ("Assumptions", "Scenario selector (one cell), 9 scenario drivers, division & product inputs, seasonality"),
        ("Hist_Data", f"Reported financials {YEARS[0]}-{YEARS[-1]}, one source per row (inputs)"),
        ("Hist_Analysis", "Growth, margins, segment mix, free cash flow, working capital (formulas)"),
        ("Ops_Model", "Product x month operating model: demand, capacity, yield, inventory, revenue, cost, gross profit"),
        ("Revenue", f"Monthly revenue forecast by product line, mix, quarters, FY vs {EN.base_year} base year"),
        ("PnL", f"Monthly P&L to net income; EBITDA; {EN.base_year} base-year comparison"),
        ("FCF", "Capex and depreciation, working capital, free cash flow"),
        ("Budget", f"{EN.fy} budget = Base scenario, locked (does not follow the selector)"),
        ("Actuals", f"SIMULATED actual results {fyc['actuals_label']}, same operating logic"),
        ("Variance", "YTD budget vs actual: price/volume/mix, operating-profit bridge, reconciliation flags"),
        ("Capex", f"Hypothetical {CFG['capex']['project']} proposal: NPV, IRR, payback, break-even, recommendation by scenario"),
        ("Dash_Export", "Tidy tables read by the dashboard export script (live scenario + scenario-independent tables)"),
        ("Checks", "Independent checks and hand-calculable unit tests; overall status feeds this cover")]
for i, (t, d) in enumerate(tabs):
    cv.cell(row=14 + i, column=2, value=t).font = f_bold
    cv.cell(row=14 + i, column=2).hyperlink = f"#'{t}'!A1"
    cv.cell(row=14 + i, column=3, value=d).font = f_base
r = 14 + len(tabs) + 1
cv.cell(row=r, column=2, value="Color legend").font = f_sec
legend = [("1,234.5", f_input, None, "Blue: hardcoded input (reported value or assumption)"),
          ("1,234.5", f_calc, None, "Black: formula / calculation"),
          ("1,234.5", f_link, None, "Green: direct link to another tab"),
          ("Input", f_base, fill_key, "Yellow fill: key assumption / cell a user may change")]
for i, (s, fnt, fill, d) in enumerate(legend):
    c = cv.cell(row=r + 1 + i, column=2, value=s)
    c.font = fnt
    if fill:
        c.fill = fill
    cv.cell(row=r + 1 + i, column=3, value=d).font = f_base
r += len(legend) + 2
cv.cell(row=r, column=2, value="Value classification").font = f_sec
for i, d in enumerate(["Reported public data - transcribed from an SEC filing (source on every row)",
                       "Calculated value - Excel formula on other cells",
                       "User-provided assumption - logged in Sources_Log with who decided it",
                       "Simulated operational data - labeled with the disclaimer above"]):
    cv.cell(row=r + 1 + i, column=3, value=d).font = f_base
r += 6
cv.cell(row=r, column=2, value="Units & conventions").font = f_sec
cv.cell(row=r, column=3, value=f"USD millions unless noted. Fiscal year ends {co['fiscal_year_end']}. "
                                "Negatives in parentheses. Capex printed negative in Hist_Data, positive in analysis.").font = f_base
cv.sheet_view.showGridLines = False

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {"Cover": NAVY, "Sources_Log": "7F7F7F", "Hist_Data": "4472C4",
                                    "Hist_Analysis": "4472C4", "Assumptions": "FFC000", "Ops_Model": "ED7D31",
                                    "Revenue": "ED7D31", "PnL": "ED7D31", "FCF": "ED7D31",
                                    "Budget": "7030A0", "Actuals": "7030A0", "Variance": "7030A0", "Capex": "C00000",
                                    "Dash_Export": "00B0F0", "Checks": "70AD47"}.get(ws.title, "FFFFFF")
ORDER = ["Cover", "Sources_Log", "Assumptions", "Hist_Data", "Hist_Analysis", "Ops_Model", "Revenue", "PnL", "FCF",
         "Budget", "Actuals", "Variance", "Capex", "Dash_Export", "Checks"]
wb._sheets = [wb[n] for n in ORDER] + [ws for ws in wb.worksheets if ws.title not in ORDER]
wb.active = 0
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Saved {OUT.relative_to(ROOT)}; Hist_Data rows={len(HD)}, Hist_Analysis metrics={len(HA)}, checks={len(CK_ROWS)}")

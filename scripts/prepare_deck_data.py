"""
Collect every number AND every sentence the executive presentation shows into outputs/<TICKER>/presentation/deck_data.json.

Numbers: from the validated dashboard export (data.json) plus a few cells read straight from the Excel workbook.
No number is typed by hand.

Sentences: generated from the numbers by explicit rules (which product is the bottleneck, whether YTD is ahead or
behind, what the capex rule concluded, ...). Any sentence can be overridden per company in the [narrative] section of
companies/<TICKER>/config.toml; overrides are templates filled with the same numbers ({ytd_op_var}, {tl}, ...), so an
override can change wording but not invent figures. Rules and overrides are listed in docs/engine.md.

Run: python scripts/prepare_deck_data.py --company TICKER   (default MCHP)
"""
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import engine  # noqa: E402

EN = engine.get()
ROOT = engine.ROOT
D = json.loads(EN.data_json.read_text())
wb = load_workbook(EN.model, data_only=True)
CFG = EN.cfg

# ---------------- formatting (same conventions as build_presentation.js) ----------------
MINUS = "−"


def f1(v):
    return f"{abs(v):,.1f}"


def money(v):
    return (MINUS if v < 0 else "") + "$" + f1(v) + "m"


def smoney(v):
    return ("+" if v >= 0 else MINUS) + "$" + f1(v) + "m"


def pct(v, d=1):
    return f"{v * 100:.{d}f}%"


def spct(v, d=1):
    return ("+" if v >= 0 else MINUS) + f"{abs(v) * 100:.{d}f}%"


def spct0(v):
    """signed percent, no decimals when whole (e.g. +5%, -4%, +1.5%)"""
    x = abs(v) * 100
    s = f"{x:.0f}" if abs(x - round(x)) < 1e-9 else f"{x:.1f}"
    return ("+" if v >= 0 else MINUS) + s + "%"


def article(num_text):
    """'an' before numbers read with a leading vowel sound (8, 11, 18, 80-89 ...)."""
    t = num_text.lstrip("$+" + MINUS)
    return "an" if t.startswith("8") or t.startswith("11") or t.startswith("18") else "a"


def join_and(xs):
    return xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " and " + xs[-1]


# ---------------- workbook cells ----------------
def row(ws, label, start=1, col=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if v == label or (isinstance(v, str) and label.endswith("*") and v.startswith(label[:-1])):
            return r
    raise KeyError(label)


hd = wb["Hist_Data"]
nY = len(EN.years)
C_LID = 6 + nY
q = None
if CFG["company"].get("quarter"):
    q_start = row(hd, "Quarterly context*")
    rq = row(hd, "net_sales", start=q_start, col=C_LID)
    q = {"prior": hd.cell(row=rq, column=3).value, "latest": hd.cell(row=rq, column=4).value}
    q["growth"] = q["latest"] / q["prior"] - 1
ck = wb["Checks"]
statuses = [ck.cell(row=r, column=9).value for r in range(1, ck.max_row + 1)]
bt = json.loads((EN.tests / "phase7_dashboard_validation.json").read_text())
asm = wb["Assumptions"]
r_prod0, r_cap0 = row(asm, "Units produced (tested)"), row(asm, CFG["division"].get("capacity_definition", "Internal final-test capacity"))
util0 = (sum(asm.cell(row=r_prod0, column=3 + i).value for i in range(3)) /
         (12 * sum(asm.cell(row=r_cap0, column=3 + i).value for i in range(3))))

# ---------------- facts the sentences are built from ----------------
B, S = D["scenarios"]["Base"], D["static"]
K = {x["metric_id"]: x["value"] for x in B["kpis"]}
KS = {s: {x["metric_id"]: x["value"] for x in D["scenarios"][s]["kpis"]} for s in D["scenarios"]}
PF = {x["product_id"]: x for x in B["product_fy"]}
VK = {(x["product_id"], x["kpi_id"]): x for x in S["variance_kpis"]}
cap = {x["key"]: x["value"] for x in S["capex_inputs"]}
cs = {x["scenario"]: x for x in S["capex_summary"]}
drv = {x["driver_id"]: x for x in S["scenario_drivers"]}
br = {x["step_id"]: x["value"] for x in S["op_bridge"]}
hist = S["history"]
prods = EN.products
tag = {p["id"]: p["tag"] for p in prods}
plural = {p["id"]: p.get("plural", p["short"]) for p in prods}
cx = CFG["capex"]
TL = cx.get("target_label", EN.target["id"])
TID = EN.target["id"]
fstart = dt.date.fromisoformat(CFG["forecast"]["start"])
alast = dt.date(fstart.year + (fstart.month + 3) // 12, (fstart.month + 3) % 12 + 1, 1)
anext = dt.date(alast.year + alast.month // 12, alast.month % 12 + 1, 1)

# bottleneck = highest Base peak demand / good-unit capacity; yield problem = biggest YTD yield shortfall;
# soft product = biggest YTD utilization cut vs budget
bottleneck = max(prods, key=lambda p: PF[p["id"]]["peak_demand_capacity_ratio"])["id"]
yld_gap = {p["id"]: VK[(p["id"], "yld")]["actual"] - VK[(p["id"], "yld")]["budget"] for p in prods}
yprob = min(yld_gap, key=yld_gap.get)
util_gap = {p["id"]: VK[(p["id"], "util")]["actual"] - VK[(p["id"], "util")]["budget"] for p in prods}
soft = min(util_gap, key=util_gap.get)
reco = cap["recommendation"]
verdict = "REJECT" if reco.startswith("REJECT") else "CHECKPOINT" if "CHECKPOINT" in reco else "APPROVE"
passing = [s for s in ("Downside", "Base", "Upside") if cs[s]["meets_all"] == "YES"]
failing = [s for s in ("Downside", "Base", "Upside") if s not in passing]
steps = {"b_price": "price", "b_vol": "sales volume", "b_mix": "product mix", "b_mat": "material cost",
         "b_yld": "manufacturing yield", "b_fix": "fixed manufacturing spending", "b_abs": "production volume & inventory",
         "b_rd": "R&D", "b_sga": "SG&A"}
top = max(steps, key=lambda k: abs(br[k]))
ytd_op = K["ytd_op_variance"]
peak = max(hist, key=lambda h: h["net_sales"])
trough = min(hist, key=lambda h: h["net_sales"])
last = hist[-1]
gm_peak = max(hist, key=lambda h: h["gross_margin"])
gm_trough = min(hist, key=lambda h: h["gross_margin"])
qc = CFG["company"].get("quarter", {})

ctx = {  # every placeholder a [narrative] override may use
    "company": EN.company["name"], "short_name": EN.company.get("short_name", EN.company["name"].split()[0]),
    "legal_name": EN.company["legal_name"], "exchange": EN.company["exchange"], "ticker": EN.ticker,
    "division": EN.division["name"], "fy": EN.fy, "fys": "FY" + EN.fy[-2:], "by": EN.base_year,
    "tl": TL, "target": TID, "project": cx["project"], "project_short": cx.get("project_short", cx["project"]),
    "capex": "$" + f1(cap["capex"]) + "m", "bottleneck": tag[bottleneck], "yprob": tag[yprob], "soft": tag[soft],
    "ytd_op_var": smoney(ytd_op), "ytd_rev_var": smoney(K["ytd_revenue_variance"]),
    "base_util": pct(drv["util"]["Base"], 0), "util0": pct(util0, 0), "act_range": f"{fstart:%b}–{alast:%b}",
    "act_long": f"{fstart:%B}–{alast:%B}", "next_month": f"{anext:%B}",
    "dn_growth": spct0(cx["growth_after"][0]), "be_util": pct(cs["Base"]["break_even_utilization"], 0),
    "peak_year": peak["fiscal_year"], "trough_year": trough["fiscal_year"], "last_year": last["fiscal_year"],
    "dem_base": spct0(drv["dem_g"]["Base"]), "q_growth": spct(q["growth"]) if q else "",
    "q_latest": qc.get("latest", ""), "q_prior": qc.get("prior", ""), "q_growth0": pct(q["growth"], 0) if q else "",
    "op_budget": "", "op_actual": "",
}

N = {}  # generated sentences
# slide 1
N["deck_title"] = f"{ctx['division']} - {EN.fy} Executive Review"
N["eyebrow_date"] = f"FP&A PORTFOLIO PROJECT  ·  {CFG['forecast']['as_of'].upper()}"
N["title_sub"] = f"{EN.fy} plan, year-to-date performance and a {ctx['capex']} capital decision"
N["title_footer"] = (f"A fictional division modeled on {ctx['legal_name']} ({ctx['exchange']}: {ctx['ticker']}). Historical figures are "
                     f"{ctx['short_name']}'s reported SEC data; all division-level data is simulated. ")
N["title_notes"] = (f"Purpose: review the division's {EN.fy} outlook, year-to-date results versus budget, the scenario range, and one capital request. "
                    "Everything comes from one linked Excel model; the dashboard and this deck read the same validated export.")
# slide 2
N["s2_title"] = f"{'Ahead of' if ytd_op >= 0 else 'Behind'} plan, with one constraint to fix: {tag[bottleneck]} test capacity"
N["s2_m1_head"] = "Ahead of budget" if ytd_op >= 0 else "Behind budget"
N["s2_mix"] = (f"Product mix {smoney(br['b_mix'])}, price {smoney(br['b_price'])} and volume {smoney(br['b_vol'])} "
               "versus budget operating profit.")
N["s2_m2_head"] = f"{tag[bottleneck]} is the bottleneck"
vb, vy = VK[(bottleneck, "yld")], VK[(bottleneck, "util")]
N["s2_m2"] = (f"{tag[bottleneck]} yield {pct(vb['actual'])} vs {pct(vb['budget'])} plan; testers already at {pct(vy['actual'], 0)} utilization, "
              f"and demand reaches {PF[bottleneck]['peak_demand_capacity_ratio']:.2f}× of full-load good output.")
verb = {"CHECKPOINT": f"Approve the {ctx['capex']} {ctx['project_short']} with a demand checkpoint",
        "APPROVE": f"Approve the {ctx['capex']} {ctx['project_short']}",
        "REJECT": f"Do not approve the {ctx['capex']} {ctx['project_short']}"}[verdict]
tail = "".join(f" It fails in {s} (NPV {money(cs[s]['npv'])})." for s in failing if s != "Base") if verdict != "REJECT" else ""
N["s2_m3"] = (f"{verb}: Base NPV {money(cs['Base']['npv'])}, IRR {pct(cs['Base']['irr']) if isinstance(cs['Base']['irr'], float) else 'n.m.'}, "
              f"payback {cs['Base']['payback_years']:.1f} years." if isinstance(cs['Base']['payback_years'], (int, float)) else
              f"{verb}: Base NPV {money(cs['Base']['npv'])}.") + tail
N["s2_notes"] = (f"Lead with the answer. Base-case {EN.fy} revenue {money(K['revenue'])}, operating profit {money(K['operating_profit'])}, free cash flow {money(K['fcf'])}. "
                 f"Year to date we are {smoney(ytd_op)} {'ahead' if ytd_op >= 0 else 'behind'} on operating profit. "
                 f"The one operational constraint is {tag[bottleneck]} final test"
                 + (", which is also the subject of the capital request." if bottleneck == TID else "."))
# slide 3
_words = {3: "three", 4: "four", 5: "five", 6: "six", 7: "seven"}
_fil = CFG["company"].get("filings_summary", "SEC filings").replace(" (SEC EDGAR)", "").replace("Forms ", "")
N["s3_reported"] = f"{ctx['company']} {_fil}: {_words.get(len(hist), len(hist))} years of history, each value traced to its SEC filing."
N["s3_sim"] = (f"Division demand, capacity, yield, unit costs and the {ctx['act_range']} {alast:%Y} actuals. "
               "Labeled as simulated on every tab and page.")
# slide 4
N["s4_eyebrow"] = f"Context · {ctx['company']} reported results"
if trough is last:
    N["s4_title"] = f"Sales fell to a {last['fiscal_year']} low" + (f"; the latest quarter is {ctx['q_growth']} year over year" if q else "")
else:
    N["s4_title"] = (f"The cycle bottomed in {trough['fiscal_year']} and demand is recovering"
                     + (" fast" if (q and q["growth"] > 0.2) else ""))
N["s4_chart"] = f"{ctx['short_name']} net sales, $m (fiscal years end {CFG['company']['fiscal_year_end'].split()[0]})"
N["s4_stat1_label"] = f"Sales change {peak['fiscal_year']} → {trough['fiscal_year']}"
N["s4_stat1_value"] = spct(trough["net_sales"] / peak["net_sales"] - 1)
N["s4_stat1_sub"] = "Peak to trough, reported net sales"
if q:
    N["s4_stat2_label"] = f"{qc['latest']} sales vs {qc['prior']}"
    N["s4_stat2_value"] = spct(q["growth"])
    N["s4_stat2_sub"] = f"${f1(q['latest'])}m vs ${f1(q['prior'])}m (10-Q)"
else:
    N["s4_stat2_label"] = f"Sales change {hist[-2]['fiscal_year']} → {last['fiscal_year']}"
    N["s4_stat2_value"] = spct(last["net_sales"] / hist[-2]["net_sales"] - 1)
    N["s4_stat2_sub"] = f"${f1(last['net_sales'])}m vs ${f1(hist[-2]['net_sales'])}m"
N["s4_stat3_label"] = (f"Gross margin {gm_peak['fiscal_year']} \u2192 {gm_trough['fiscal_year']}" if gm_trough is last else
                      f"Gross margin {gm_peak['fiscal_year']} \u2192 {gm_trough['fiscal_year']} \u2192 {last['fiscal_year']}")
N["s4_stat3_value"] = pct(gm_trough["gross_margin"])
N["s4_stat3_sub"] = (f"down from the {pct(gm_peak['gross_margin'])} peak" if gm_trough is last else
                    f"from {pct(gm_peak['gross_margin'])} peak; {pct(last['gross_margin'])} in {last['fiscal_year']}")
N["s4_notes"] = (f"Reported public data only on this slide. Net sales moved {spct(trough['net_sales'] / peak['net_sales'] - 1)} from the "
                 f"{peak['fiscal_year']} peak to the {trough['fiscal_year']} trough; gross margin ranged from {pct(gm_trough['gross_margin'])} to "
                 f"{pct(gm_peak['gross_margin'])}. The Base demand-growth assumption is {ctx['dem_base']}.")
# slide 5
N["s5_notes"] = (f"Operating profit grows faster than revenue because fixed factory cost is spread over more units (utilization {ctx['util0']} → {ctx['base_util']}). "
                 f"The outlook line replaces {ctx['act_long']} with actuals; the remaining months are forecast at budget in the Base case.")
# slide 6
N["s6_eyebrow"] = f"Budget vs actual · YTD {ctx['act_long']} {alast:%Y}"
N["s6_title"] = (f"Operating profit is {smoney(ytd_op)} {'ahead of' if ytd_op >= 0 else 'behind'} budget, "
                 f"{'driven by' if (br[top] >= 0) == (ytd_op >= 0) else 'despite'} {steps[top]}")
vp = {x["line_id"]: x for x in S["variance_pnl"]}
ctx["op_budget"], ctx["op_actual"] = money(vp["op"]["budget"]), money(vp["op"]["actual"])
fav = [steps[k] for k in steps if br[k] > 0.005]
unf = [steps[k] for k in steps if br[k] < -0.005]
N["s6_notes"] = (f"Budget YTD operating profit {money(vp['op']['budget'])}, actual {money(vp['op']['actual'])}. "
                 f"The largest driver is {steps[top]} ({smoney(br[top])}). Favorable: {join_and(fav) if fav else 'none'}; "
                 f"unfavorable: {join_and(unf) if unf else 'none'}. "
                 "Every component is computed separately and reconciles to the total with zero difference.")
# slide 7
N["s7_title"] = f"{tag[bottleneck]} runs closest to its test-capacity limit, peaking at {PF[bottleneck]['peak_demand_capacity_ratio']:.2f}"
vyp = VK[(yprob, "yld")]
N["tiles"] = [
    [f"{tag[yprob]} yield, YTD", pct(vyp["actual"]), f"plan {pct(vyp['budget'])}" + (f" \u00b7 {cx['yield_tile_note']}" if cx.get("yield_tile_note") else ""), "coral"],
    [f"{tag[bottleneck]} tester utilization, YTD", pct(vy["actual"], 0), f"plan {pct(vy['budget'], 0)} · limited headroom", "warn"],
    [f"{tag[yprob]} cost per good unit, YTD", "$" + f"{VK[(yprob, 'cpgu')]['actual']:.2f}",
     f"plan ${VK[(yprob, 'cpgu')]['budget']:.2f} · yield loss adds cost" if yld_gap[yprob] < 0 else f"plan ${VK[(yprob, 'cpgu')]['budget']:.2f}", "ink"],
    [f"{tag[soft]} utilization, YTD", pct(VK[(soft, 'util')]["actual"], 0),
     "cut to avoid building inventory" if util_gap[soft] < 0 else f"plan {pct(VK[(soft, 'util')]['budget'], 0)}", "ink"]]
N["s7_notes"] = (f"The heat map compares each month's demand with what the test floor can deliver at full load. {tag[bottleneck]} runs closest to its limit, "
                 f"peaking at {PF[bottleneck]['peak_demand_capacity_ratio']:.2f} in the Base plan. Year to date {tag[yprob]} yield is {pct(vyp['actual'])} "
                 f"against {article(pct(vyp['budget']))} {pct(vyp['budget'])} plan, which {'raises' if yld_gap[yprob] < 0 else 'lowers'} cost per good unit.")
# slide 8
d_ = lambda s: (f"demand {spct0(drv['dem_g'][s])}, price {spct0(drv['asp_chg'][s])}"  # noqa: E731
                + (f", yield {spct0(drv['yld_adj'][s]).replace('%', ' pts')}" if drv["yld_adj"][s] < 0 else "")
                + f", {pct(drv['util'][s], 0)} utilization")
N["s8_notes"] = (f"All three cases run through the same Excel model via one selector cell. Downside: {d_('Downside')}. "
                 f"Upside: {d_('Upside')}. FY outlook vs budget ranges from {smoney(KS['Downside']['outlook_op_vs_budget'])} "
                 f"to {smoney(KS['Upside']['outlook_op_vs_budget'])}.")
# slide 9
noun = cx.get("project_noun", "The project")
if len(passing) == 3:
    N["s9_title"] = f"{noun} passes every hurdle in all three scenarios"
elif not passing:
    N["s9_title"] = f"{noun} fails the hurdles in every scenario"
else:
    N["s9_title"] = f"{noun} passes every hurdle in {join_and(passing)}, not {join_and(failing)}"
N["s9_eyebrow"] = f"Capital allocation · hypothetical {cx['project']}"
N["s9_notes"] = (f"Break-even needs about {pct(cs['Base']['break_even_utilization'], 0)} use of the added capacity in Base. "
                 + (f"In Downside ({TL} demand {ctx['dn_growth']} a year) the NPV is {money(cs['Downside']['npv'])}, "
                    f"so the {ctx['capex']} is not repaid at the {pct(cap['rate'], 0)} discount rate."
                    if cs["Downside"]["npv"] < 0 else
                    f"Even in Downside the NPV is {money(cs['Downside']['npv'])}."))
# slide 10
N["s10_title"] = {"CHECKPOINT": f"Approve the {ctx['project_short']}, released on a demand checkpoint",
                  "APPROVE": f"Approve the {ctx['project_short']}",
                  "REJECT": f"Do not approve the {ctx['project_short']} as proposed"}[verdict]
N["s10_steps"] = [
    [f"Track {tag[yprob]} yield monthly",
     f"Recover from {pct(vyp['actual'])} toward the {pct(vyp['budget'])} plan; every point of yield lowers cost per good unit."
     if yld_gap[yprob] < 0 else f"Hold yield at {pct(vyp['actual'])} or better; every point of yield lowers cost per good unit."],
    ["Re-forecast at the Q2 close", f"Load {ctx['next_month']} actuals and re-run all three scenarios before the next capital review."],
    [f"Hold {tag[soft].lower() if tag[soft][1:].islower() else tag[soft]} loading",
     f"Keep utilization near {pct(VK[(soft, 'util')]['actual'], 0)} while {tag[soft].lower() if tag[soft][1:].islower() else tag[soft]} demand trails plan, to avoid building inventory."],
    ["Firm up the proposal", "Replace the hypothetical price, capacity and yield gain with vendor quotes, then re-test the decision rule."]]
N["s10_notes"] = ("Decision rule: reject if Base fails; approve outright only if Downside also passes; otherwise approve with a checkpoint. "
                  f"Base NPV {money(cs['Base']['npv'])}, IRR {pct(cs['Base']['irr']) if isinstance(cs['Base']['irr'], float) else 'n.m.'}, "
                  f"payback {cs['Base']['payback_years']:.2f} years; " if isinstance(cs['Base']['payback_years'], (int, float)) else
                  "Decision rule: reject if Base fails; approve outright only if Downside also passes; otherwise approve with a checkpoint. "
                  f"Base NPV {money(cs['Base']['npv'])}; ") + f"Downside NPV {money(cs['Downside']['npv'])}."
N["act_range"] = ctx["act_range"]
N["outlook_series"] = f"Outlook (actual {ctx['act_range']} + forecast)"

# ---------------- per-company wording overrides ----------------
ov = CFG.get("narrative", {})
for k, v in ov.items():
    if k not in N:
        raise SystemExit(f"[narrative] {k}: unknown key (valid: {', '.join(sorted(N))})")
    fill = lambda t: t.format(**ctx) if isinstance(t, str) else [fill(x) for x in t]  # noqa: E731
    N[k] = fill(v)

scen = {}
for s in ["Downside", "Base", "Upside"]:
    k = {x["metric_id"]: x["value"] for x in D["scenarios"][s]["kpis"]}
    scen[s] = {"kpis": k, "product_fy": D["scenarios"][s]["product_fy"], "monthly": D["scenarios"][s]["monthly"],
               "product_monthly": D["scenarios"][s]["product_monthly"]}
out = {
    "company": D["company"],
    "disclaimer": D["disclaimer"],
    "source_sha256": D["source_sha256"],
    "scenarios": scen,
    "static": D["static"],
    "q_prior_sales": q and q["prior"], "q_latest_sales": q and q["latest"], "q_growth": q and q["growth"],
    "base_year_utilization": util0,
    "excel_checks_pass": statuses.count("PASS"), "excel_checks_fail": statuses.count("FAIL"),
    "dashboard_test": {"compared": bt["compared"], "failed": bt["failed"], "total": bt["total_checks"],
                       "matches_this_export": bt["data_sha256"] == D["source_sha256"]},
    "narrative": N,
    "deck_file": EN.deck.name,
}
EN.deck_data.write_text(json.dumps(out, indent=1, default=str))
print(f"Wrote {EN.deck_data.relative_to(ROOT)}; checks {out['excel_checks_pass']}/{out['excel_checks_pass'] + out['excel_checks_fail']}, "
      f"bottleneck {bottleneck}, verdict {verdict}, overrides {len(ov)}, dashboard test matches export: {out['dashboard_test']['matches_this_export']}")

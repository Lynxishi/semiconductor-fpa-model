# Phase 2 - Assumptions (scenario controls) and Ops_Model (product x month).
# Executed inside build_model.py's namespace (shares styles and helpers). Not a standalone script.
import datetime as _dt

from openpyxl.worksheet.datavalidation import DataValidation

SIM = "Simulated operational data"
ASSUME = "User-provided assumption"
CALC = "Calculated value"
f_sim = Font(name=F, size=10, color="0000FF")  # simulated inputs are hardcoded inputs -> blue

PRODUCTS = [(p["key"], p["name"]) for p in EN.products]
PCOL = {p["key"]: c for p, c in zip(EN.products, "CDE")}  # product columns on Assumptions
BY, FYL = EN.base_year, EN.fy
FSTART = _dt.datetime.fromisoformat(CFG["forecast"]["start"])
_mlab = lambda d: d.strftime("%b-%Y")  # noqa: E731
FEND = _dt.datetime(FSTART.year + (FSTART.month + 10) // 12, (FSTART.month + 10) % 12 + 1, 1)
PN = CFG.get("product_notes", {})
MONTH_COLS = [get_column_letter(3 + i) for i in range(12)]  # C..N
FYC = "O"  # forecast-year total column on Ops_Model

ASM = {}   # key -> row on Assumptions
OPS = {}   # (product, key) -> row on Ops_Model


def build_phase2_tabs():
    global r
    # =================================================================
    # Assumptions
    # =================================================================
    asm = wb.create_sheet("Assumptions")
    title_block(asm, "Assumptions and Scenario Controls",
                "Change ONLY the yellow selector (C6) to switch scenario. Blue = input. Every forecast tab reads the Live column.")
    setw(asm, {"A": 58, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12,
               "J": 12, "K": 12, "L": 12, "M": 12, "N": 12, "O": 12, "P": 26, "Q": 70})
    asm.freeze_panes = "B5"

    # --- scenario selector
    asm["A6"] = "ACTIVE SCENARIO  (choose Downside / Base / Upside)"
    asm["A6"].font = Font(name=F, size=11, bold=True, color=NAVY)
    asm["C6"] = "Base"
    asm["C6"].font = Font(name=F, size=11, bold=True, color="0000FF")
    asm["C6"].fill = fill_key
    asm["C6"].alignment = Alignment(horizontal="center")
    dv = DataValidation(type="list", formula1='"Downside,Base,Upside"', allow_blank=False,
                        showErrorMessage=True, errorTitle="Scenario", error="Choose Downside, Base or Upside")
    asm.add_data_validation(dv)
    dv.add("C6")
    asm["D6"] = "=MATCH($C$6,$C$9:$E$9,0)"
    asm["D6"].font = f_calc
    asm["E6"] = "<- scenario index (1 = Downside, 2 = Base, 3 = Upside)"
    asm["E6"].font = f_sub
    ASM["sel"], ASM["idx"] = 6, 6

    # --- scenario drivers
    r = 8
    section(asm, r, f"A. Scenario drivers ({FYL} vs {BY} base year)", 17)
    r += 1
    header_row(asm, r, ["Driver", "Units", "Downside", "Base", "Upside", "LIVE"] + [""] * 9 + ["Classification", "Rationale / reference"])
    ASM["drv_hdr"] = r
    r += 1
    DLAB = {"dem_g": ("Customer demand growth", "%"), "asp_chg": ("Average selling price change", "%"),
            "yld_adj": ("Manufacturing yield change (percentage points vs base-year yield)", "pts"),
            "util": ("Capacity utilization (internal final test)", "%"),
            "mat_inf": ("Material-cost inflation (variable cost per unit)", "%"), "rd_g": ("R&D expense growth", "%"),
            "sga_g": ("SG&A expense growth", "%"),
            "capex_pct": ("Capital expenditures (% of revenue, excl. proposed test-equipment project)", "%"),
            "tax": ("Tax rate", "%")}
    DRIVERS = [(k, DLAB[k][0], DLAB[k][1], *CFG["drivers"][k]["values"], CFG["drivers"][k].get("why", ""))
               for k in engine.DRIVER_KEYS]
    ASM["drv_first"] = r
    for key, lab, units, d, b, u, why in DRIVERS:
        ASM[key] = r
        asm.cell(row=r, column=1, value=lab).font = f_base
        asm.cell(row=r, column=2, value=units).font = f_base
        for col, v in zip("CDE", (d, b, u)):
            c = asm[f"{col}{r}"]
            c.value, c.font, c.number_format = v, f_input, PCT
        c = asm[f"F{r}"]
        c.value = f"=INDEX(C{r}:E{r},1,$D$6)"
        c.font, c.number_format, c.fill = f_bold, PCT, PatternFill("solid", fgColor="FFF2CC")
        asm.cell(row=r, column=16, value=ASSUME).font = f_base
        asm.cell(row=r, column=17, value=why).font = f_base
        r += 1
    ASM["drv_last"] = r - 1

    # --- historical reference points (formulas on Hist tabs)
    r += 1
    section(asm, r, f"B. Reference points from {COMPANY}'s reported history (formulas - context for the drivers above)", 17)
    r += 1
    PY = YEARS[-2]
    peak_gm_year = max(YEARS, key=lambda y: float(REP[(y, "gross_profit")]["value_usd_m"]) / float(REP[(y, "net_sales")]["value_usd_m"]))
    REFS = []
    if QROW:
        REFS += [(f"{QCFG['latest']} net sales growth vs {QCFG['prior']}", f"=Hist_Analysis!$E${HA['q:net_sales']}", PCT),
                 (f"{QCFG['latest']} sales x 4 vs {BY} net sales", f"=Hist_Data!$D${QROW['net_sales']}*4/{H('net_sales', BY)}-1", PCT)]
    REFS += [
        (f"Gross margin {BY} (reported basis)", f"={AA('gm', BY)}", PCT),
        (f"Gross margin {peak_gm_year} (highest in the history)", f"={AA('gm', peak_gm_year)}", PCT),
        (f"R&D growth {BY}", f"={H('rd', BY)}/{H('rd', PY)}-1", PCT),
        (f"SG&A growth {BY}", f"={H('sga', BY)}/{H('sga', PY)}-1", PCT),
        (f"Capex % of sales {BY}", f"={AA('capex_pct', BY)}", PCT),
        (f"Capex % of sales, average {YEARS[0]}-{YEARS[-1]}",
         f"=AVERAGE(Hist_Analysis!${YCOL[YEARS[0]]}${HA['capex_pct']}:${YCOL[YEARS[-1]]}${HA['capex_pct']})", PCT),
        (f"Effective tax rate, average {YEARS[0]}-{YEARS[-1]} (years shown n.m. excluded)",
         f"=AVERAGE(Hist_Analysis!${YCOL[YEARS[0]]}${HA['etr']}:${YCOL[YEARS[-1]]}${HA['etr']})", PCT),
    ]
    if "fact:internal_test_share" in FACT_ROW:
        REFS.append((f"Share of test performed internally ({BY})", f"=Hist_Data!${YCOL[BY]}${FACT_ROW['fact:internal_test_share']}", '0%'))
    for lab, formula, fmt in REFS:
        asm.cell(row=r, column=1, value=lab).font = f_base
        c = asm[f"F{r}"]
        c.value, c.font, c.number_format = formula, f_link, fmt
        asm.cell(row=r, column=16, value=CALC).font = f_base
        asm.cell(row=r, column=17, value="Reported public data (Hist tabs)").font = f_base
        r += 1

    # --- division base year
    r += 1
    section(asm, r, f"C. Fictional {EN.division['name']} - base year {BY} (" + DISCLAIMER + ")", 17)
    r += 1
    ASM["base_rev"] = r
    asm.cell(row=r, column=1, value=f"Division net revenue, {BY} base year").font = f_bold
    asm.cell(row=r, column=2, value="$m").font = f_base
    c = asm[f"F{r}"]
    c.value, c.font, c.number_format, c.fill = float(EN.division["base_revenue"]), f_input, USD, fill_key
    asm.cell(row=r, column=16, value=ASSUME).font = f_base
    asm.cell(row=r, column=17, value=f"User-provided assumption. ~{EN.division['base_revenue'] / float(REP[(BY, 'net_sales')]['value_usd_m']):.0%} of {COMPANY} {BY} net sales").font = f_base
    r += 1
    for key, lab, val, why in [("rd0", f"Division R&D expense, {BY} base year", float(EN.division["rd0"]), EN.division.get("rd_note", "")),
                               ("sga0", f"Division SG&A expense, {BY} base year", float(EN.division["sga0"]), EN.division.get("sga_note", ""))]:
        ASM[key] = r
        asm.cell(row=r, column=1, value=lab).font = f_base
        asm.cell(row=r, column=2, value="$m").font = f_base
        c = asm[f"F{r}"]
        c.value, c.font, c.number_format = val, f_sim, USD
        asm.cell(row=r, column=16, value=SIM).font = f_base
        asm.cell(row=r, column=17, value=why + " Used in Phase 3.").font = f_base
        r += 1
    ASM["fstart"] = r
    asm.cell(row=r, column=1, value=f"Forecast start month (first month of {FYL})").font = f_base
    c = asm[f"F{r}"]
    c.value, c.font, c.number_format = FSTART, f_input, "mmm-yy"
    asm.cell(row=r, column=16, value=ASSUME).font = f_base
    asm.cell(row=r, column=17, value=f"{FYL} = {_mlab(FSTART)} to {_mlab(FEND)}, calendar months").font = f_base
    r += 1

    # --- product inputs
    r += 1
    section(asm, r, f"D. Product inputs - base year {BY} (simulated) and live {FYL} drivers (calculated)", 17)
    r += 1
    header_row(asm, r, ["Input", "Units"] + [p["short"] for p in EN.products] + ["Division"] + [""] * 9 + ["Classification", "Notes"])
    r += 1

    def prow(key, label, units, vals=None, fn=None, fmt=USD, cls=SIM, note="", div=None, bold=False):
        """vals: dict product->hardcoded value; fn: product->formula (no '='). div: formula for Division column."""
        global r
        ASM[key] = r
        asm.cell(row=r, column=1, value=label).font = f_bold if bold else f_base
        asm.cell(row=r, column=2, value=units).font = f_base
        for p, _ in PRODUCTS:
            c = asm[f"{PCOL[p]}{r}"]
            if vals is not None:
                c.value, c.font = vals[p], f_sim if cls == SIM else f_input
            else:
                c.value, c.font = "=" + fn(p), f_calc
            c.number_format = fmt
        if div:
            c = asm[f"F{r}"]
            c.value, c.font, c.number_format = "=" + div, f_calc, fmt
        asm.cell(row=r, column=16, value=cls).font = f_base
        asm.cell(row=r, column=17, value=note).font = f_base
        r += 1

    P = lambda key, p: f"${PCOL[p]}${ASM[key]}"   # noqa: E731  product cell on Assumptions
    LIVE = lambda key: f"$F${ASM[key]}"           # noqa: E731
    MIXROW = {p["key"]: "mix:" + p["segment"] for p in EN.products}
    PV = lambda f: {p["key"]: (p.get("segment_share", 1.0) if f == "segment_share_v" else p[f]) for p in EN.products}  # noqa: E731
    seg_notes = " ".join(p["segment_note"] for p in EN.products if p.get("segment_note"))

    # A product may be a user-assumed share of one reported segment (segment_share), and the products may not cover
    # every reported segment. Then the division mix is each product's share of the reported total it covers.
    SPLIT = any(p.get("segment_share", 1.0) != 1.0 for p in EN.products) or \
        {p["segment"] for p in EN.products} != set(SEGMENTS)
    if SPLIT:
        prow("seg_share", "Share of the reported segment assigned to this product", "%", vals=PV("segment_share_v"),
             fmt=PCT, cls=ASSUME, note=CFG.get("product_notes", {}).get("segment_share",
             "User-provided assumption: split of a reported segment that the company does not report separately."))
        prow("mix_rep", "Reported segment share of net sales x assigned share", "%",
             fn=lambda p: f"{AA(MIXROW[p], BY)}*{P('seg_share', p)}", fmt='0.0000%', cls=CALC,
             note=f"{COMPANY} {BY} reported segment mix x the split above. " + seg_notes, div=f"SUM(C{r}:E{r})")
        prow("mix0", "Base-year revenue mix", "%", fn=lambda p: f"{P('mix_rep', p)}/$F${ASM['mix_rep']}", fmt='0.0000%',
             cls=CALC, note="Renormalized to 100% over the three products (reported segments not modeled are excluded).",
             div=f"SUM(C{r}:E{r})")
    else:
        prow("mix0", "Base-year revenue mix", "%", fn=lambda p: AA(MIXROW[p], BY), fmt='0.0000%', cls=CALC,
             note=f"{COMPANY} {BY} reported segment mix. " + seg_notes,
             div=f"SUM(C{r}:E{r})")
    prow("rev0", "Base-year revenue", "$m", fn=lambda p: f"$F${ASM['base_rev']}*{P('mix0', p)}", cls=CALC,
         note="Division revenue x mix", div=f"SUM(C{r}:E{r})", bold=True)
    prow("asp0", "Base-year average selling price", "$/unit", vals=PV("asp0"),
         fmt='$#,##0.00', note=PN.get("asp0", "Simulated."))
    prow("yld0", "Base-year final-test yield", "%", vals=PV("yld0"), fmt=PCT,
         note=PN.get("yld0", "Simulated. Good units / units tested."))
    prow("vc0", "Base-year variable cost per unit produced", "$/unit", vals=PV("vc0"),
         fmt='$#,##0.00', note=PN.get("vc0", "Simulated."))
    prow("fix0", "Fixed manufacturing cost per month", "$m/month", vals=PV("fix0"),
         note=PN.get("fix0", "Simulated."),
         div=f"SUM(C{r}:E{r})")
    prow("cap0", "Internal final-test capacity", "m units/month", vals=PV("cap0"),
         fmt='#,##0.00', note=PN.get("cap0", "Simulated. Capacity = units the internal test floor can process per month."))
    prow("gdiff", "Product demand growth vs division driver", "pts", vals=PV("gdiff"),
         fmt=PCT, note=PN.get("gdiff", "Simulated."))
    prow("inv_m", "Opening finished-goods inventory (months of base-year sales)", "months",
         vals=PV("inv_m"), fmt='0.00', note=PN.get("inv_m", "Simulated opening balance."))

    r += 1
    asm.cell(row=r, column=1, value=f"Base-year {BY} economics (calculated; assumes production = sales in the base year)").font = f_sec
    r += 1
    prow("sold0", "Units sold", "m units", fn=lambda p: f"{P('rev0', p)}/{P('asp0', p)}", fmt='#,##0.00', cls=CALC,
         note="Revenue / ASP")
    prow("prod0", "Units produced (tested)", "m units", fn=lambda p: f"{P('sold0', p)}/{P('yld0', p)}", fmt='#,##0.00',
         cls=CALC, note="Units sold / yield (no inventory change in base year)")
    prow("util0", "Implied capacity utilization", "%", fn=lambda p: f"{P('prod0', p)}/({P('cap0', p)}*12)", fmt=PCT,
         cls=CALC, note="Units produced / (capacity x 12).")
    prow("vcost0", "Variable manufacturing cost", "$m", fn=lambda p: f"{P('prod0', p)}*{P('vc0', p)}", cls=CALC,
         note="Units produced x variable cost per unit", div=f"SUM(C{r}:E{r})")
    prow("fcost0", "Fixed manufacturing cost", "$m", fn=lambda p: f"{P('fix0', p)}*12", cls=CALC,
         note="Monthly fixed cost x 12", div=f"SUM(C{r}:E{r})")
    prow("cogs0", "Cost of revenue", "$m", fn=lambda p: f"{P('vcost0', p)}+{P('fcost0', p)}", cls=CALC,
         note="Variable + fixed manufacturing cost", div=f"SUM(C{r}:E{r})")
    prow("gp0", "Gross profit", "$m", fn=lambda p: f"{P('rev0', p)}-{P('cogs0', p)}", cls=CALC,
         note="Revenue - cost of revenue", div=f"SUM(C{r}:E{r})", bold=True)
    prow("gm0", "Gross margin", "%", fn=lambda p: f"IFERROR({P('gp0', p)}/{P('rev0', p)},0)", fmt=PCT, cls=CALC,
         note=f"Calibration: compare with {COMPANY} {BY} reported gross margin in section B",
         div=f"IFERROR(F{ASM['gp0']}/F{ASM['rev0']},0)")
    prow("cpgu0", "Cost per good unit", "$/unit", fn=lambda p: f"{P('cogs0', p)}/{P('sold0', p)}", fmt='$#,##0.000',
         cls=CALC, note="Cost of revenue / good units (= units sold in base year). Values the opening inventory.")

    r += 1
    asm.cell(row=r, column=1, value=f"Live {FYL} product drivers (calculated from the active scenario)").font = f_sec
    r += 1
    prow("dem27", f"{FYL} customer demand", "m units",
         fn=lambda p: f"{P('sold0', p)}*(1+{LIVE('dem_g')}+{P('gdiff', p)})", fmt='#,##0.00', cls=CALC,
         note="Base-year units sold x (1 + demand growth + product differential)")
    prow("asp27", f"{FYL} average selling price", "$/unit", fn=lambda p: f"{P('asp0', p)}*(1+{LIVE('asp_chg')})",
         fmt='$#,##0.000', cls=CALC, note="Base ASP x (1 + ASP change)")
    prow("yld27", f"{FYL} final-test yield", "%", fn=lambda p: f"MIN(1,MAX(0,{P('yld0', p)}+{LIVE('yld_adj')}))",
         fmt=PCT, cls=CALC, note="Base yield + yield change, bounded 0-100%")
    prow("util27", f"{FYL} capacity utilization", "%", fn=lambda p: LIVE("util"), fmt=PCT, cls=CALC,
         note="Same scenario utilization for every product line")
    prow("vc27", f"{FYL} variable cost per unit produced", "$/unit", fn=lambda p: f"{P('vc0', p)}*(1+{LIVE('mat_inf')})",
         fmt='$#,##0.000', cls=CALC, note="Base variable cost x (1 + material inflation)")
    prow("inv_u0", f"Opening finished-goods inventory ({FSTART.strftime('%d-%b-%Y')})", "m units",
         fn=lambda p: f"{P('inv_m', p)}*{P('sold0', p)}/12", fmt='#,##0.00', cls=CALC,
         note="Months of inventory x base-year monthly units sold")
    prow("inv_v0", "Opening finished-goods inventory value", "$m", fn=lambda p: f"{P('inv_u0', p)}*{P('cpgu0', p)}",
         cls=CALC, note="Opening units x base-year cost per good unit", div=f"SUM(C{r}:E{r})")

    # --- seasonality
    r += 1
    section(asm, r, "E. Monthly demand seasonality (simulated; must sum to 12.00)", 17)
    r += 1
    ASM["mon_hdr"] = r
    asm.cell(row=r, column=1, value="Month").font = f_bold
    for i, col in enumerate(MONTH_COLS):
        c = asm[f"{col}{r}"]
        c.value = f"=DATE(YEAR($F${ASM['fstart']}),MONTH($F${ASM['fstart']})+{i},1)"
        c.number_format, c.font = "mmm-yy", f_bold
        c.alignment = Alignment(horizontal="right")
    asm[f"O{r}"] = "Sum"
    asm[f"O{r}"].font = f_bold
    r += 1
    ASM["season"] = r
    asm.cell(row=r, column=1, value="Seasonality factor (1.00 = average month)").font = f_base
    SEAS = CFG["seasonality"]["factors"]
    for col, v in zip(MONTH_COLS, SEAS):
        c = asm[f"{col}{r}"]
        c.value, c.font, c.number_format = v, f_sim, '0.00'
    asm[f"O{r}"] = f"=SUM(C{r}:N{r})"
    asm[f"O{r}"].number_format = '0.00'
    asm.cell(row=r, column=16, value=SIM).font = f_base
    asm.cell(row=r, column=17, value=CFG["seasonality"].get("note", "")).font = f_base
    r += 1

    # =================================================================
    # Ops_Model
    # =================================================================
    ops = wb.create_sheet("Ops_Model")
    title_block(ops, f"Product Operating Model - {FYL} monthly (active scenario)",
                "Units in millions, money in $m, prices in $/unit. Every row traces to Assumptions. Black = formula.")
    setw(ops, {"A": 50, "B": 12, **{c: 10 for c in MONTH_COLS}, "O": 12, "P": 24, "Q": 66})
    ops.freeze_panes = "C7"
    ops["A4"] = "Scenario shown:"
    ops["A4"].font = f_bold
    ops["B4"] = "=Assumptions!$C$6"
    ops["B4"].font = Font(name=F, size=10, bold=True, color="008000")
    ops["B4"].fill = PatternFill("solid", fgColor="FFF2CC")
    header_row(ops, 6, ["Line item", "Units"] + [""] * 12 + [FYL, "Classification", "Formula"])
    for i, col in enumerate(MONTH_COLS):
        c = ops[f"{col}6"]
        c.value = f"=Assumptions!{col}${ASM['mon_hdr']}"
        c.number_format = "mmm-yy"
        c.font, c.fill = f_hdr, fill_hdr
        c.alignment = Alignment(horizontal="center")
    OPS["month_hdr"] = 6
    r = 7

    def orow(p, key, label, units, month_fn, fy, fmt=USD, note="", bold=False, cls=None):
        """month_fn(col, prev_col_or_None) -> formula (no '='). fy: formula for FY column (no '=') or None."""
        global r
        OPS[(p, key)] = r
        ops.cell(row=r, column=1, value=label).font = f_bold if bold else f_base
        ops.cell(row=r, column=2, value=units).font = f_base
        for i, col in enumerate(MONTH_COLS):
            c = ops[f"{col}{r}"]
            c.value = "=" + month_fn(col, MONTH_COLS[i - 1] if i else None)
            c.font, c.number_format = f_calc, fmt
        if fy:
            c = ops[f"{FYC}{r}"]
            c.value, c.font, c.number_format = "=" + fy, f_bold, fmt
        ops.cell(row=r, column=16, value=cls or f"{CALC} ({SIM[:9].lower()} inputs)").font = f_base
        ops.cell(row=r, column=17, value=note).font = f_base
        if bold:
            for col in range(1, 16):
                ops.cell(row=r, column=col).border = b_top
        r += 1

    SUMFY = lambda key, p: f"SUM(C{OPS[(p, key)]}:N{OPS[(p, key)]})"  # noqa: E731

    for p, pname in PRODUCTS:
        section(ops, r, f"{pname}  -  {DISCLAIMER}", 17)
        r += 1
        A_ = lambda key: f"Assumptions!{P(key, p)}"  # noqa: E731
        R_ = lambda key, col: f"{col}{OPS[(p, key)]}"  # noqa: E731
        orow(p, "season", "Seasonality factor", "x", lambda c, w: f"Assumptions!{c}${ASM['season']}",
             f"SUM(C{r}:N{r})", fmt='0.00', note="Link: Assumptions section E")
        orow(p, "dem", "Customer demand", "m units", lambda c, w: f"{A_('dem27')}*{R_('season', c)}/12",
             f"SUM(C{r}:N{r})", fmt='#,##0.00', note=f"{FYL} demand x seasonality / 12")
        orow(p, "cap", "Internal final-test capacity", "m units", lambda c, w: A_("cap0"),
             f"SUM(C{r}:N{r})", fmt='#,##0.00', note="Monthly capacity (Assumptions D)")
        orow(p, "util", "Capacity utilization", "%", lambda c, w: A_("util27"),
             f"IFERROR({FYC}{r + 1}/{FYC}{r - 1},0)", fmt=PCT, note="Scenario utilization. FY = units produced / capacity")
        orow(p, "prod", "Units produced (tested)", "m units", lambda c, w: f"{R_('cap', c)}*{R_('util', c)}",
             f"SUM(C{r}:N{r})", fmt='#,##0.00', note="Capacity x utilization")
        orow(p, "yld", "Manufacturing yield (final test)", "%", lambda c, w: A_("yld27"),
             f"IFERROR({FYC}{r + 1}/{FYC}{r - 1},0)", fmt=PCT, note="Scenario yield. FY = good units / units produced")
        orow(p, "good", "Good units produced", "m units", lambda c, w: f"{R_('prod', c)}*{R_('yld', c)}",
             f"SUM(C{r}:N{r})", fmt='#,##0.00', note="Units produced x yield")
        orow(p, "inv_open", "Opening finished-goods inventory", "m units",
             lambda c, w: A_("inv_u0") if w is None else f"{w}{r + 3}", f"C{r}", fmt='#,##0.00',
             note="First month: Assumptions opening balance; then prior month's closing")
        orow(p, "avail", "Available good units", "m units", lambda c, w: f"{R_('inv_open', c)}+{R_('good', c)}",
             None, fmt='#,##0.00', note="Opening inventory + good units produced")
        orow(p, "sold", "Units sold", "m units", lambda c, w: f"MIN({R_('dem', c)},{R_('avail', c)})",
             f"SUM(C{r}:N{r})", fmt='#,##0.00', note="MIN(customer demand, available good units)", bold=True)
        orow(p, "inv_close", "Closing finished-goods inventory", "m units",
             lambda c, w: f"{R_('avail', c)}-{R_('sold', c)}", f"N{r}", fmt='#,##0.00', note="Available - units sold")
        orow(p, "unmet", "Unmet demand (supply shortfall)", "m units", lambda c, w: f"{R_('dem', c)}-{R_('sold', c)}",
             f"SUM(C{r}:N{r})", fmt='#,##0.00', note="Demand - units sold. >0 means lost or delayed sales")
        orow(p, "dcr", "Demand / good-unit capacity (risk ratio)", "x",
             lambda c, w: f"IFERROR({R_('dem', c)}/({R_('cap', c)}*{R_('yld', c)}),0)",
             f"IFERROR({FYC}{OPS[(p, 'dem')]}/({FYC}{OPS[(p, 'cap')]}*{FYC}{OPS[(p, 'yld')]}),0)", fmt='0.00',
             note="Demand / (capacity x yield) = share of full-load good output needed. Feeds the risk heat map")
        orow(p, "asp", "Average selling price", "$/unit", lambda c, w: A_("asp27"),
             f"IFERROR({FYC}{r + 1}/{FYC}{OPS[(p, 'sold')]},0)", fmt='$#,##0.000', note="Scenario ASP. FY = revenue / units sold")
        orow(p, "rev", "Revenue", "$m", lambda c, w: f"{R_('sold', c)}*{R_('asp', c)}", f"SUM(C{r}:N{r})",
             note="Units sold x ASP", bold=True)
        orow(p, "vcu", "Variable cost per unit produced", "$/unit", lambda c, w: A_("vc27"),
             f"IFERROR({FYC}{r + 1}/{FYC}{OPS[(p, 'prod')]},0)", fmt='$#,##0.000', note="Scenario variable cost per unit")
        orow(p, "vcost", "Variable manufacturing cost", "$m", lambda c, w: f"{R_('prod', c)}*{R_('vcu', c)}",
             f"SUM(C{r}:N{r})", note="Units produced x variable cost per unit")
        orow(p, "fcost", "Fixed manufacturing cost", "$m", lambda c, w: A_("fix0"), f"SUM(C{r}:N{r})",
             note="Monthly fixed cost (Assumptions D)")
        orow(p, "mfg", "Total manufacturing cost", "$m", lambda c, w: f"{R_('vcost', c)}+{R_('fcost', c)}",
             f"SUM(C{r}:N{r})", note="Variable + fixed manufacturing cost")
        orow(p, "cpgu", "Cost per good unit", "$/unit", lambda c, w: f"IFERROR({R_('mfg', c)}/{R_('good', c)},0)",
             f"IFERROR({FYC}{r - 1}/{FYC}{OPS[(p, 'good')]},0)", fmt='$#,##0.000',
             note="Total manufacturing cost / good units. Rises when yield or utilization falls")
        orow(p, "invv_open", "Opening inventory value", "$m",
             lambda c, w: A_("inv_v0") if w is None else f"{w}{r + 3}", f"C{r}", note="Prior month's closing value")
        orow(p, "avgc", "Average cost per available unit", "$/unit",
             lambda c, w: f"IFERROR(({R_('invv_open', c)}+{R_('mfg', c)})/{R_('avail', c)},0)", None, fmt='$#,##0.000',
             note="(Opening value + manufacturing cost) / available units - weighted-average costing")
        orow(p, "cor", "Cost of revenue", "$m", lambda c, w: f"{R_('sold', c)}*{R_('avgc', c)}", f"SUM(C{r}:N{r})",
             note="Units sold x average cost")
        orow(p, "invv_close", "Closing inventory value", "$m",
             lambda c, w: f"{R_('invv_open', c)}+{R_('mfg', c)}-{R_('cor', c)}", f"N{r}",
             note="Opening value + manufacturing cost - cost of revenue")
        orow(p, "cm", "Contribution margin", "$m",
             lambda c, w: f"{R_('rev', c)}-{R_('sold', c)}*IFERROR({R_('vcu', c)}/{R_('yld', c)},0)", f"SUM(C{r}:N{r})",
             note="Revenue - units sold x variable cost per good unit (variable cost / yield)")
        orow(p, "cm_pct", "Contribution margin %", "%", lambda c, w: f"IFERROR({R_('cm', c)}/{R_('rev', c)},0)",
             f"IFERROR({FYC}{r - 1}/{FYC}{OPS[(p, 'rev')]},0)", fmt=PCT, note="Contribution margin / revenue")
        orow(p, "gp", "Gross profit", "$m", lambda c, w: f"{R_('rev', c)}-{R_('cor', c)}", f"SUM(C{r}:N{r})",
             note="Revenue - cost of revenue", bold=True)
        orow(p, "gm", "Gross margin %", "%", lambda c, w: f"IFERROR({R_('gp', c)}/{R_('rev', c)},0)",
             f"IFERROR({FYC}{r - 1}/{FYC}{OPS[(p, 'rev')]},0)", fmt=PCT, note="Gross profit / revenue")
        r += 1

    # --- division total
    section(ops, r, f"Division total  -  {DISCLAIMER}", 17)
    r += 1
    SUMP = lambda key: (lambda c, w: "+".join(f"{c}{OPS[(p, key)]}" for p, _ in PRODUCTS))  # noqa: E731
    for key, lab, bold in [("rev", "Revenue", True), ("vcost", "Variable manufacturing cost", False),
                           ("fcost", "Fixed manufacturing cost", False), ("mfg", "Total manufacturing cost", False),
                           ("cor", "Cost of revenue", False), ("cm", "Contribution margin", False)]:
        orow("div", key, lab, "$m", SUMP(key), f"SUM(C{r}:N{r})", bold=bold, note="Sum of the three product lines")
    orow("div", "cm_pct", "Contribution margin %", "%", lambda c, w: f"IFERROR({c}{OPS[('div', 'cm')]}/{c}{OPS[('div', 'rev')]},0)",
         f"IFERROR({FYC}{OPS[('div', 'cm')]}/{FYC}{OPS[('div', 'rev')]},0)", fmt=PCT, note="Contribution margin / revenue")
    orow("div", "gp", "Gross profit", "$m", SUMP("gp"), f"SUM(C{r}:N{r})", bold=True, note="Sum of the three product lines")
    orow("div", "gm", "Gross margin %", "%", lambda c, w: f"IFERROR({c}{OPS[('div', 'gp')]}/{c}{OPS[('div', 'rev')]},0)",
         f"IFERROR({FYC}{OPS[('div', 'gp')]}/{FYC}{OPS[('div', 'rev')]},0)", fmt=PCT, note="Gross profit / revenue")
    orow("div", "invv_close", "Closing finished-goods inventory value", "$m", SUMP("invv_close"),
         f"N{r}", note="Sum of product closing values. Feeds working capital in Phase 3")
    orow("div", "lost_rev", "Revenue at risk from unmet demand", "$m",
         lambda c, w: "+".join(f"{c}{OPS[(p, 'unmet')]}*{c}{OPS[(p, 'asp')]}" for p, _ in PRODUCTS),
         f"SUM(C{r}:N{r})", note="Unmet units x ASP, summed across products")

    # Scenario summary back on Assumptions (live scenario only)
    r_s = asm.max_row + 2
    section(asm, r_s, f"F. Live scenario result - {FYL} (from Ops_Model)", 17)
    items = [("Division revenue", f"=Ops_Model!${FYC}${OPS[('div', 'rev')]}", USD),
             (f"Growth vs {BY} base-year revenue", f"=Ops_Model!${FYC}${OPS[('div', 'rev')]}/$F${ASM['base_rev']}-1", PCT),
             ("Gross profit", f"=Ops_Model!${FYC}${OPS[('div', 'gp')]}", USD),
             ("Gross margin", f"=Ops_Model!${FYC}${OPS[('div', 'gm')]}", PCT),
             ("Revenue at risk from unmet demand", f"=Ops_Model!${FYC}${OPS[('div', 'lost_rev')]}", USD)]
    for i, (lab, fml, fmt) in enumerate(items):
        rr = r_s + 1 + i
        asm.cell(row=rr, column=1, value=lab).font = f_base
        c = asm[f"F{rr}"]
        c.value, c.font, c.number_format = fml, f_link, fmt
        asm.cell(row=rr, column=16, value=CALC).font = f_base
    ASM["summary"] = r_s + 1


def add_phase2_checks():
    """Appends Phase 2 checks to the Checks tab. Result in column C, tolerance H, status I."""
    global r
    cksec("Phase 2 - scenario controls and operating model (result = absolute difference or violation count)")
    header_row(ck, r, ["Check", "Area", "Result", "", "", "", "", "Tolerance", "Status"])
    r += 1
    rows_ = [p for p, _ in PRODUCTS]
    M = lambda p, key: f"Ops_Model!$C${OPS[(p, key)]}:$N${OPS[(p, key)]}"   # noqa: E731
    FYV = lambda p, key: f"Ops_Model!${FYC}${OPS[(p, key)]}"                 # noqa: E731
    AS = lambda key, col: f"Assumptions!${col}${ASM[key]}"                 # noqa: E731
    d0, d1 = ASM["drv_first"], ASM["drv_last"]

    def one(label, area, formula, tol=0.0001):
        global r
        ck.cell(row=r, column=1, value=label).font = f_base
        ck.cell(row=r, column=2, value=area).font = f_base
        c = ck[f"C{r}"]
        c.value, c.font, c.number_format = "=" + formula, f_calc, '0.0000;(0.0000);"-"'
        t = ck[f"H{r}"]
        t.value, t.font = tol, f_input
        s = ck[f"I{r}"]
        s.value, s.font = f'=IF(ABS(C{r})<=H{r},"PASS","FAIL")', f_bold
        CK_ROWS.append(r)
        r += 1

    one("Scenario selector holds a valid scenario name", "Scenario",
        f'IF(ISNUMBER(Assumptions!$D$6),0,1)', 0)
    one("Live column = selected scenario column (all 9 drivers)", "Scenario",
        f"SUMPRODUCT(ABS(Assumptions!$F${d0}:$F${d1}-((Assumptions!$D$6=1)*Assumptions!$C${d0}:$C${d1}"
        f"+(Assumptions!$D$6=2)*Assumptions!$D${d0}:$D${d1}+(Assumptions!$D$6=3)*Assumptions!$E${d0}:$E${d1})))")
    one("Ops_Model shows the selected scenario name", "Scenario", 'IF(Ops_Model!$B$4=Assumptions!$C$6,0,1)', 0)
    one("Seasonality factors sum to 12.00", "Dates & units", f"Assumptions!$O${ASM['season']}-12")
    one(f"12 consecutive months, {_mlab(FSTART)} to {_mlab(FEND)}", "Dates & units",
        f"(Ops_Model!$C$6<>DATE({FSTART.year},{FSTART.month},1))+(Ops_Model!$N$6<>DATE({FEND.year},{FEND.month},1))"
        f"+SUMPRODUCT(--(DAY(Ops_Model!$C$6:$N$6)<>1))+(COUNT(Ops_Model!$C$6:$N$6)<>12)"
        f"+SUMPRODUCT((Ops_Model!$D$6:$N$6-Ops_Model!$C$6:$M$6<28)+(Ops_Model!$D$6:$N$6-Ops_Model!$C$6:$M$6>31))", 0)
    one("Base-year mix sums to 100%", "Revenue", f"{AS('mix0', 'F')}-1")
    one("Base-year product revenue sums to division revenue", "Revenue", f"{AS('rev0', 'F')}-{AS('base_rev', 'F')}", 0.05)
    one("Utilization and yield within 0-100% (violations)", "Operations",
        "+".join(f"SUMPRODUCT(({M(p, 'util')}<0)+({M(p, 'util')}>1)+({M(p, 'yld')}<0)+({M(p, 'yld')}>1))" for p in rows_), 0)
    one("Units produced = capacity x utilization (all months, all products)", "Operations",
        "+".join(f"SUMPRODUCT(ABS({M(p, 'prod')}-{M(p, 'cap')}*{M(p, 'util')}))" for p in rows_))
    one("Good units = units produced x yield", "Operations",
        "+".join(f"SUMPRODUCT(ABS({M(p, 'good')}-{M(p, 'prod')}*{M(p, 'yld')}))" for p in rows_))
    one("Units sold = MIN(demand, available good units)", "Operations",
        "+".join(f"SUMPRODUCT(ABS({M(p, 'sold')}-(({M(p, 'dem')}<{M(p, 'avail')})*{M(p, 'dem')}"
                 f"+({M(p, 'dem')}>={M(p, 'avail')})*{M(p, 'avail')})))" for p in rows_))
    one("No negative inventory (violations)", "Operations",
        "+".join(f"SUMPRODUCT(--({M(p, 'inv_close')}<-0.000001))+SUMPRODUCT(--({M(p, 'invv_close')}<-0.000001))" for p in rows_), 0)
    one("Inventory units roll forward: opening + good - sold = closing (FY)", "Operations",
        "+".join(f"ABS({FYV(p, 'inv_open')}+{FYV(p, 'good')}-{FYV(p, 'sold')}-{FYV(p, 'inv_close')})" for p in rows_))
    one("Inventory value rolls forward: opening + mfg cost - cost of revenue = closing (FY)", "Operations",
        "+".join(f"ABS({FYV(p, 'invv_open')}+{FYV(p, 'mfg')}-{FYV(p, 'cor')}-{FYV(p, 'invv_close')})" for p in rows_))
    one("Revenue = SUMPRODUCT(units sold, ASP) vs FY total (monthly-to-annual)", "Revenue",
        "+".join(f"ABS(SUMPRODUCT({M(p, 'sold')},{M(p, 'asp')})-{FYV(p, 'rev')})" for p in rows_))
    one("Variable cost = SUMPRODUCT(units produced, cost per unit) vs FY total", "Cost",
        "+".join(f"ABS(SUMPRODUCT({M(p, 'prod')},{M(p, 'vcu')})-{FYV(p, 'vcost')})" for p in rows_))
    one("Division revenue = sum of product FY revenue", "Revenue",
        f"{FYV('div', 'rev')}-(" + "+".join(FYV(p, "rev") for p in rows_) + ")")
    one("Division gross profit = revenue - cost of revenue", "Gross profit",
        f"{FYV('div', 'gp')}-({FYV('div', 'rev')}-{FYV('div', 'cor')})")
    one("Division monthly revenue sums to FY (independent SUM of product rows)", "Revenue",
        f"SUM(Ops_Model!$C${OPS[('div', 'rev')]}:$N${OPS[('div', 'rev')]})-(" +
        "+".join(f"SUM({M(p, 'rev')})" for p in rows_) + ")")
    one("Contribution margin >= gross profit (fixed cost absorbed in cost of revenue)", "Gross profit",
        f"MAX(0,{FYV('div', 'gp')}-{FYV('div', 'cm')})")

    header_row(ck, r, ["Unit test (hand-calculable)", "", "Result", "Expected", "Difference", "", "", "Tolerance", "Status"])
    r += 1
    tests2 = [
        ("Capacity 100 x utilization 80% = 80 units produced", "=100*0.8", 80),
        ("80 produced x 90% yield = 72 good units", "=80*0.9", 72),
        ("Demand 70, available 0 + 72 -> units sold = MIN(70,72) = 70", "=MIN(70,0+72)", 70),
        ("Demand 90, available 72 -> sold 72, unmet 18", "=90-MIN(90,72)", 18),
        ("70 units x ASP $2.00 = $140 revenue", "=70*2", 140),
        ("Weighted avg: open 10u @ $1 + 72u costing $90 -> $100/82 per unit", "=(10+90)/82", 100 / 82),
        ("Cost of revenue: 70 sold x ($100/82) = $85.366", "=70*(10+90)/82", 7000 / 82),
        ("Closing value: 10 + 90 - 85.366 = 12 units x 1.2195", "=10+90-70*(10+90)/82", 12 * 100 / 82),
        ("Contribution: $140 - 70 x ($0.45 / 90% yield) = $105", "=140-70*(0.45/0.9)", 105),
    ]
    for lab, formula, exp in tests2:
        ck.cell(row=r, column=1, value=lab).font = f_base
        ck[f"C{r}"] = formula
        ck[f"D{r}"] = exp
        ck[f"D{r}"].font = f_input
        ck[f"E{r}"] = f"=C{r}-D{r}"
        ck[f"H{r}"] = 0.000001
        ck[f"H{r}"].font = f_input
        ck[f"I{r}"] = f'=IF(ABS(E{r})<=H{r},"PASS","FAIL")'
        ck[f"I{r}"].font = f_bold
        CK_ROWS.append(r)
        r += 1
